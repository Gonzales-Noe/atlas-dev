"""
Syllabus template: exact-file preview, click map, and DOCX fill-for-export.

Flow:
  1. Upload DOCX / PDF / Excel — store the original file as the clickable source
  2. Extract native click targets (paragraphs, PDF text blocks, sheet cells)
  3. Keep a working DOCX copy for faculty export fill (identity for DOCX uploads)
  4. Admin clicks the real document → sections recorded; place content regions
"""

from __future__ import annotations

import hashlib
import logging
import re
from io import BytesIO

logger = logging.getLogger(__name__)

_HEADING_STYLE_RE = re.compile(r'^Heading\s+(\d+)$', re.I)
_MD_HEADING_RE = re.compile(r'^(#{1,6})\s+(.*\S)\s*$')
_TABLE_SEP_RE = re.compile(r'^\|?\s*:?-{3,}')

SUPPORTED_UPLOAD_TYPES = {'docx', 'pdf', 'xlsx', 'xls'}
MIME_BY_KIND = {
    'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    'pdf': 'application/pdf',
    'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'xls': 'application/vnd.ms-excel',
}


def _read_bytes(file_obj):
    if hasattr(file_obj, 'seek'):
        file_obj.seek(0)
    if isinstance(file_obj, (bytes, bytearray)):
        return bytes(file_obj)
    return file_obj.read()


def _docx_from_markdown(markdown):
    """Build a simple DOCX from Markdown-ish text (headings + pipe tables)."""
    from docx import Document

    doc = Document()
    lines = (markdown or '').splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        heading = _MD_HEADING_RE.match(line.strip())
        if heading:
            level = min(len(heading.group(1)), 3)
            doc.add_heading(heading.group(2).strip(), level=level)
            i += 1
            continue
        if line.strip().startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                raw = lines[i].strip()
                i += 1
                if _TABLE_SEP_RE.match(raw):
                    continue
                cells = [c.strip() for c in raw.strip('|').split('|')]
                if any(cells):
                    rows.append(cells)
            if rows:
                width = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=width)
                table.style = 'Table Grid'
                for r_i, row in enumerate(rows):
                    for c_i in range(width):
                        table.rows[r_i].cells[c_i].text = row[c_i] if c_i < len(row) else ''
            continue
        # Bold-only line as heading-ish
        bold = re.match(r'^\*\*(.+?)\*\*:?$', line.strip())
        if bold and len(bold.group(1).split()) <= 12:
            doc.add_heading(bold.group(1).strip(), level=2)
            i += 1
            continue
        doc.add_paragraph(line.strip())
        i += 1

    if not doc.paragraphs and not doc.tables:
        doc.add_paragraph('(Empty document after conversion)')
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def docx_from_pdf(file_obj):
    """Convert PDF → editable DOCX via fast Markdown extract (LiteParse / text)."""
    from .document_processor import DocumentProcessor

    data = _read_bytes(file_obj)
    pages = []
    try:
        pages = DocumentProcessor.extract_markdown_from_pdf(BytesIO(data))
    except Exception:
        logger.warning('LiteParse PDF→MD failed; trying plain text', exc_info=True)
    if not pages:
        try:
            pages = DocumentProcessor.extract_text(BytesIO(data), 'pdf')
        except Exception:
            logger.exception('Plain PDF text extract failed')
            pages = []
    markdown = '\n\n'.join(text for text, _page in pages if text and str(text).strip())
    if not markdown.strip():
        raise ValueError(
            'Could not extract text from that PDF. Use a text-based PDF, or upload a .docx / .xlsx instead.'
        )
    return _docx_from_markdown(markdown)


def docx_from_xlsx(file_obj):
    """Convert Excel workbook → editable DOCX (one heading + table per sheet)."""
    import openpyxl
    from docx import Document

    data = _read_bytes(file_obj)
    try:
        wb = openpyxl.load_workbook(BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError(f'Could not read that Excel file: {exc}') from exc

    doc = Document()
    sheet_count = 0
    for sheet in wb.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = ['' if c is None else str(c).strip() for c in row]
            if any(cells):
                rows.append(cells)
        if not rows:
            continue
        sheet_count += 1
        title = (sheet.title or f'Sheet {sheet_count}').strip()
        doc.add_heading(title, level=1)
        width = max(len(r) for r in rows)
        # Cap huge sheets for the click UI
        shown = rows[:80]
        table = doc.add_table(rows=len(shown), cols=width)
        table.style = 'Table Grid'
        for r_i, row in enumerate(shown):
            for c_i in range(width):
                table.rows[r_i].cells[c_i].text = row[c_i] if c_i < len(row) else ''
        if len(rows) > len(shown):
            doc.add_paragraph(f'… {len(rows) - len(shown)} more rows omitted from preview')

    wb.close()
    if sheet_count == 0:
        raise ValueError('That Excel file has no readable cells. Add content or upload a .docx instead.')
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def normalize_upload_to_docx(file_obj, file_type):
    """
    Build a working DOCX for export fill. Returns (docx_bytes, conversion_note).
    conversion_note is empty when the upload was already DOCX.
    """
    kind = (file_type or '').lower().lstrip('.')
    if kind not in SUPPORTED_UPLOAD_TYPES:
        raise ValueError(
            f'Unsupported file type “{kind}”. Upload a .docx, .pdf, or .xlsx syllabus sample.'
        )
    if kind == 'docx':
        return _read_bytes(file_obj), ''
    if kind == 'pdf':
        docx_bytes = docx_from_pdf(file_obj)
        return docx_bytes, 'Working DOCX derived from PDF for export fill.'
    docx_bytes = docx_from_xlsx(file_obj)
    return docx_bytes, 'Working DOCX derived from Excel for export fill.'


def _attach_working_region_ids(blocks, working_blocks):
    """Match native preview blocks to working-DOCX paragraph/table ids by text hash."""
    by_hash = {}
    for wb in working_blocks or []:
        h = wb.get('text_hash')
        if h and h not in by_hash:
            by_hash[h] = wb.get('id') or ''
    for block in blocks or []:
        h = block.get('text_hash')
        if h and h in by_hash:
            block['working_region_id'] = by_hash[h]
    return blocks


def build_exact_preview(file_obj, file_type):
    """
    Keep the original bytes for clickable preview; also build a working DOCX.

    Returns dict:
      original_bytes, kind, blocks, block_count, working_docx_bytes, note
    """
    kind = (file_type or '').lower().lstrip('.')
    if kind not in SUPPORTED_UPLOAD_TYPES:
        raise ValueError(
            f'Unsupported file type “{kind}”. Upload a .docx, .pdf, or .xlsx syllabus sample.'
        )
    original = _read_bytes(file_obj)
    if kind == 'docx':
        blocks = extract_docx_preview(BytesIO(original))['blocks']
        working = original
        note = ''
    elif kind == 'pdf':
        blocks = extract_pdf_preview(BytesIO(original))['blocks']
        working, note = normalize_upload_to_docx(BytesIO(original), 'pdf')
        working_blocks = extract_docx_preview(BytesIO(working))['blocks']
        _attach_working_region_ids(blocks, working_blocks)
    else:
        blocks = extract_xlsx_preview(BytesIO(original))['blocks']
        working, note = normalize_upload_to_docx(BytesIO(original), kind)
        working_blocks = extract_docx_preview(BytesIO(working))['blocks']
        _attach_working_region_ids(blocks, working_blocks)

    if not blocks:
        raise ValueError('No readable content found in that file.')

    return {
        'original_bytes': original,
        'kind': kind,
        'blocks': blocks,
        'block_count': len(blocks),
        'working_docx_bytes': working,
        'note': note,
        'mime': MIME_BY_KIND.get(kind, 'application/octet-stream'),
    }


def _slug(text):
    return re.sub(r'[^a-z0-9]+', '-', str(text or '').lower()).strip('-') or 'section'


def _text_hash(text):
    return hashlib.sha1((text or '').encode('utf-8')).hexdigest()[:12]


def _unique_key(base, used):
    key = base or f'section-{len(used) + 1}'
    candidate, n = key, 2
    while candidate in used:
        candidate, n = f'{key}-{n}', n + 1
    used.add(candidate)
    return candidate


def extract_docx_preview(file_obj):
    """
    Return { blocks: [...], template_name hint only from caller }.

    Each block:
      id, kind (heading|paragraph|table), level, text, style,
      rows (for tables), text_hash
    """
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph as DocxParagraph

    if hasattr(file_obj, 'seek'):
        file_obj.seek(0)
    data = file_obj.read() if not isinstance(file_obj, (bytes, bytearray)) else file_obj
    doc = Document(BytesIO(data))

    blocks = []
    para_index = 0
    table_index = 0

    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit('}', 1)[-1]
        if tag == 'p':
            paragraph = DocxParagraph(child, doc)
            text = (paragraph.text or '').strip()
            try:
                style = (paragraph.style.name if paragraph.style else '') or ''
            except Exception:
                style = ''
            heading = _HEADING_STYLE_RE.match(style)
            kind = 'heading' if heading or style == 'Title' else 'paragraph'
            level = int(heading.group(1)) if heading else (1 if style == 'Title' else 0)
            # Skip empty paragraphs (still assign ids only for non-empty to keep clicks useful)
            if not text:
                para_index += 1
                continue
            block_id = f'p-{para_index}'
            blocks.append({
                'id': block_id,
                'kind': kind,
                'level': level,
                'text': text,
                'style': style,
                'text_hash': _text_hash(text),
                'anchor': {'type': 'paragraph', 'index': para_index, 'text_hash': _text_hash(text)},
            })
            para_index += 1
        elif tag == 'tbl':
            table = DocxTable(child, doc)
            rows = []
            for row in table.rows:
                try:
                    cells = row.cells
                except Exception:
                    continue
                rows.append([' '.join(cell.text.split()) for cell in cells])
            rows = [r for r in rows if any(r)]
            preview = ' | '.join(rows[0]) if rows else f'Table {table_index + 1}'
            block_id = f't-{table_index}'
            blocks.append({
                'id': block_id,
                'kind': 'table',
                'level': 0,
                'text': preview[:200],
                'style': 'Table',
                'rows': rows[:40],
                'text_hash': _text_hash(preview),
                'anchor': {'type': 'table', 'index': table_index, 'text_hash': _text_hash(preview)},
            })
            table_index += 1

    # Keep qn imported for future oxml tweaks without unused-import noise in some linters
    _ = qn
    return {'blocks': blocks, 'block_count': len(blocks)}


def extract_pdf_preview(file_obj):
    """Native PDF text blocks with page + bbox for overlay clicks on the real PDF."""
    try:
        import pymupdf
    except ImportError:
        import fitz as pymupdf  # type: ignore

    data = _read_bytes(file_obj)
    doc = pymupdf.open(stream=data, filetype='pdf')
    blocks = []
    try:
        for page_index, page in enumerate(doc):
            page_dict = page.get_text('dict') or {}
            block_i = 0
            for raw in page_dict.get('blocks') or []:
                if raw.get('type') != 0:
                    continue
                lines = []
                for line in raw.get('lines') or []:
                    span_text = ''.join(span.get('text') or '' for span in (line.get('spans') or []))
                    if span_text.strip():
                        lines.append(span_text.strip())
                text = ' '.join(lines).strip()
                if not text:
                    block_i += 1
                    continue
                bbox = raw.get('bbox') or [0, 0, 0, 0]
                block_id = f'pdf-{page_index}-b-{block_i}'
                blocks.append({
                    'id': block_id,
                    'kind': 'paragraph',
                    'level': 0,
                    'text': text[:2000],
                    'style': '',
                    'text_hash': _text_hash(text),
                    'page': page_index,
                    'bbox': [float(bbox[0]), float(bbox[1]), float(bbox[2]), float(bbox[3])],
                    'page_width': float(page.rect.width),
                    'page_height': float(page.rect.height),
                    'anchor': {
                        'type': 'pdf_block',
                        'page': page_index,
                        'index': block_i,
                        'text_hash': _text_hash(text),
                    },
                })
                block_i += 1
    finally:
        doc.close()
    return {'blocks': blocks, 'block_count': len(blocks)}


def extract_xlsx_preview(file_obj):
    """Native Excel cells — same grid the UI renders for click mapping."""
    import openpyxl

    data = _read_bytes(file_obj)
    try:
        wb = openpyxl.load_workbook(BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError(f'Could not read that Excel file: {exc}') from exc

    blocks = []
    sheets = []
    try:
        for sheet_index, sheet in enumerate(wb.worksheets):
            title = (sheet.title or f'Sheet {sheet_index + 1}').strip()
            sheet_id = f'xls-{sheet_index}-title'
            blocks.append({
                'id': sheet_id,
                'kind': 'heading',
                'level': 1,
                'text': title,
                'style': 'Sheet',
                'text_hash': _text_hash(title),
                'sheet': sheet_index,
                'sheet_name': title,
                'row': None,
                'col': None,
                'anchor': {'type': 'sheet_title', 'sheet': sheet_index, 'text_hash': _text_hash(title)},
            })
            max_row = 0
            max_col = 0
            cells = []
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if value is None or str(value).strip() == '':
                        continue
                    text = str(value).strip()
                    r = cell.row
                    c = cell.column
                    max_row = max(max_row, r)
                    max_col = max(max_col, c)
                    if r > 80 or c > 40:
                        continue
                    cell_id = f'xls-{sheet_index}-r{r}-c{c}'
                    block = {
                        'id': cell_id,
                        'kind': 'paragraph',
                        'level': 0,
                        'text': text[:500],
                        'style': 'Cell',
                        'text_hash': _text_hash(text),
                        'sheet': sheet_index,
                        'sheet_name': title,
                        'row': r,
                        'col': c,
                        'anchor': {
                            'type': 'cell',
                            'sheet': sheet_index,
                            'row': r,
                            'col': c,
                            'text_hash': _text_hash(text),
                        },
                    }
                    blocks.append(block)
                    cells.append({'id': cell_id, 'row': r, 'col': c, 'text': text[:500]})
            sheets.append({
                'index': sheet_index,
                'name': title,
                'max_row': min(max_row, 80) or 1,
                'max_col': min(max_col, 40) or 1,
                'cells': cells,
            })
    finally:
        wb.close()

    if not any(b.get('kind') != 'heading' for b in blocks):
        raise ValueError('That Excel file has no readable cells.')
    return {'blocks': blocks, 'block_count': len(blocks), 'sheets': sheets}


def field_map_to_syllabus_sections(field_map):
    """Derive faculty-compatible syllabus_sections from the click map."""
    used = set()
    sections = []
    items = (field_map or {}).get('sections') if isinstance(field_map, dict) else field_map
    if not isinstance(items, list):
        return []

    for index, item in enumerate(items):
        if not isinstance(item, dict):
            continue
        label = str(item.get('label') or '').strip()
        if not label:
            continue
        key = item.get('key') or _unique_key(_slug(label), used)
        if key not in used:
            used.add(key)
        sub_used = set()
        subsections = []
        for sub in item.get('subsections') or []:
            if not isinstance(sub, dict):
                continue
            sub_label = str(sub.get('label') or '').strip()
            if not sub_label:
                continue
            sub_key = sub.get('key') or _unique_key(_slug(sub_label), sub_used)
            subsections.append({
                'key': sub_key,
                'label': sub_label,
                'required': bool(sub.get('required')),
                'table_enabled': bool(sub.get('table_enabled')),
                'table_columns': list(sub.get('table_columns') or []),
                'table_column_groups': list(sub.get('table_column_groups') or []),
                'title_region_id': sub.get('title_region_id') or '',
                'content_region_id': sub.get('content_region_id') or '',
            })
        sections.append({
            'key': key,
            'label': label,
            'enabled': item.get('enabled', True) is not False,
            'is_static': bool(item.get('is_static')),
            'static_content': str(item.get('static_content') or ''),
            'table_enabled': bool(item.get('table_enabled')),
            'table_columns': list(item.get('table_columns') or []),
            'table_column_groups': list(item.get('table_column_groups') or []),
            'subsections': subsections,
            'title_region_id': item.get('title_region_id') or '',
            'content_region_id': item.get('content_region_id') or '',
            'editor': item.get('editor') or None,
            'confidence': item.get('confidence'),
        })
        # Built-in editor key override
        editor = item.get('editor') or item.get('key')
        if editor in {'basic', 'details', 'iloat', 'schedule', 'tos', 'grading', 'policies'}:
            sections[-1]['key'] = editor
    return sections


def _iter_body_paragraphs(doc):
    from docx.text.paragraph import Paragraph as DocxParagraph
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit('}', 1)[-1]
        if tag == 'p':
            yield DocxParagraph(child, doc)


def _iter_body_tables(doc):
    from docx.table import Table as DocxTable
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit('}', 1)[-1]
        if tag == 'tbl':
            yield DocxTable(child, doc)


def _set_paragraph_text(paragraph, text):
    text = '' if text is None else str(text)
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        paragraph.add_run(text)


def _fill_region_id(region_id, blocks_by_id=None):
    """Prefer working DOCX region id when the click was on PDF/Excel."""
    if not region_id:
        return ''
    rid = str(region_id)
    if rid.startswith('p-') or rid.startswith('t-'):
        return rid
    if blocks_by_id and rid in blocks_by_id:
        working = blocks_by_id[rid].get('working_region_id') or ''
        if working:
            return working
    return rid


def _resolve_paragraph(doc, region_id, blocks_by_id=None):
    rid = _fill_region_id(region_id, blocks_by_id)
    if not rid or not rid.startswith('p-'):
        return None
    try:
        index = int(rid.split('-', 1)[1])
    except ValueError:
        return None
    for i, paragraph in enumerate(_iter_body_paragraphs(doc)):
        if i == index:
            return paragraph
    return None


def _resolve_table(doc, region_id, blocks_by_id=None):
    rid = _fill_region_id(region_id, blocks_by_id)
    if not rid or not rid.startswith('t-'):
        return None
    try:
        index = int(rid.split('-', 1)[1])
    except ValueError:
        return None
    for i, table in enumerate(_iter_body_tables(doc)):
        if i == index:
            return table
    return None


def _answer_for_section(syllabus, section):
    """Best-effort text for a mapped section from faculty syllabus data."""
    key = str(section.get('key') or '')
    label = str(section.get('label') or '')
    custom = syllabus.get('custom_subsections') if isinstance(syllabus.get('custom_subsections'), dict) else {}
    section_answers = custom.get(key) if isinstance(custom.get(key), dict) else {}

    # Concatenate subsection answers when present
    parts = []
    for sub in section.get('subsections') or []:
        sub_key = str(sub.get('key') or '')
        answer = str(section_answers.get(sub_key) or '').strip()
        if answer:
            parts.append(answer)
    if parts:
        return '\n\n'.join(parts)

    # Built-in / top-level fields by key or common aliases
    aliases = {
        'basic': ['course_title', 'course_code', 'course_description'],
        'details': ['course_description', 'rationale', 'teaching_strategies'],
        'policies': ['policies'],
        'grading': ['grading_system'],
        'vision': ['vision'],
        'mission': ['mission'],
    }
    if key in aliases:
        chunks = [str(syllabus.get(a) or '').strip() for a in aliases[key]]
        chunks = [c for c in chunks if c]
        if chunks:
            return '\n\n'.join(chunks)

    for candidate in (key, _slug(label), label):
        val = syllabus.get(candidate)
        if isinstance(val, str) and val.strip():
            return val.strip()
        if isinstance(val, list) and val:
            return '\n'.join(str(x) for x in val if x)

    if section.get('is_static') and section.get('static_content'):
        return str(section['static_content'])
    return ''


def fill_docx_from_field_map(template_bytes, field_map, syllabus, preview_blocks=None):
    """
    Clone working DOCX bytes, write answers into content region targets.
    PDF/Excel clicks resolve via working_region_id on preview blocks.
    """
    from docx import Document

    doc = Document(BytesIO(template_bytes))
    sections = (field_map or {}).get('sections') if isinstance(field_map, dict) else []
    if not isinstance(sections, list):
        sections = []

    blocks_by_id = {}
    for block in preview_blocks or []:
        if isinstance(block, dict) and block.get('id'):
            blocks_by_id[block['id']] = block
    # Also allow working ids stored directly on mapped sections
    for section in sections:
        if not isinstance(section, dict):
            continue
        for key in ('title_region_id', 'content_region_id', 'working_content_region_id'):
            rid = section.get(key)
            if rid and rid not in blocks_by_id and str(rid).startswith(('p-', 't-')):
                blocks_by_id[rid] = {'id': rid, 'working_region_id': rid}
        for sub in section.get('subsections') or []:
            if not isinstance(sub, dict):
                continue
            for key in ('title_region_id', 'content_region_id', 'working_content_region_id'):
                rid = sub.get(key)
                if rid and rid not in blocks_by_id and str(rid).startswith(('p-', 't-')):
                    blocks_by_id[rid] = {'id': rid, 'working_region_id': rid}

    custom = syllabus.get('custom_subsections') if isinstance(syllabus.get('custom_subsections'), dict) else {}
    custom_tables = syllabus.get('custom_tables') if isinstance(syllabus.get('custom_tables'), dict) else {}

    for section in sections:
        if not isinstance(section, dict) or section.get('enabled') is False:
            continue
        section_key = str(section.get('key') or '')

        section_answers = custom.get(section_key) if isinstance(custom.get(section_key), dict) else {}
        section_tables = custom_tables.get(section_key) if isinstance(custom_tables.get(section_key), dict) else {}
        for sub in section.get('subsections') or []:
            if not isinstance(sub, dict):
                continue
            sub_key = str(sub.get('key') or '')
            content_id = sub.get('working_content_region_id') or sub.get('content_region_id') or ''
            if not content_id:
                continue
            fill_id = _fill_region_id(content_id, blocks_by_id)
            if sub.get('table_enabled') and fill_id.startswith('t-'):
                table = _resolve_table(doc, fill_id, blocks_by_id)
                rows = section_tables.get(sub_key) or []
                columns = sub.get('table_columns') or []
                if table is not None and rows:
                    _fill_table(table, columns, rows)
            else:
                paragraph = _resolve_paragraph(doc, fill_id, blocks_by_id)
                answer = str(section_answers.get(sub_key) or '').strip()
                if paragraph is not None and answer:
                    _set_paragraph_text(paragraph, answer)

        content_id = section.get('working_content_region_id') or section.get('content_region_id') or ''
        fill_id = _fill_region_id(content_id, blocks_by_id)
        if fill_id.startswith('p-'):
            paragraph = _resolve_paragraph(doc, fill_id, blocks_by_id)
            answer = _answer_for_section(syllabus, section)
            if paragraph is not None and answer and not section.get('is_static'):
                _set_paragraph_text(paragraph, answer)
        elif fill_id.startswith('t-'):
            table = _resolve_table(doc, fill_id, blocks_by_id)
            rows = (custom_tables.get(section_key) or {}).get('__section__') or []
            if table is not None and rows:
                _fill_table(table, section.get('table_columns') or [], rows)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _fill_table(table, columns, rows):
    """Write data rows under an existing header row when possible."""
    headers = [str(c or '').strip() for c in (columns or [])]
    start_row = 1 if table.rows else 0
    # Ensure enough rows
    needed = start_row + len(rows)
    while len(table.rows) < needed:
        table.add_row()
    col_count = len(table.columns) if table.columns else (len(headers) or 1)
    for offset, row_values in enumerate(rows):
        if not isinstance(row_values, list):
            continue
        row = table.rows[start_row + offset]
        for col in range(min(col_count, len(row.cells))):
            value = row_values[col] if col < len(row_values) else ''
            row.cells[col].text = '' if value is None else str(value)


def normalize_field_map(field_map):
    """Sanitize map from the admin UI before save."""
    if not isinstance(field_map, dict):
        return {'sections': []}
    sections = []
    used = set()
    for item in field_map.get('sections') or []:
        if not isinstance(item, dict):
            continue
        label = str(item.get('label') or '').strip()
        if not label:
            continue
        key = _unique_key(str(item.get('key') or _slug(label)), used)
        subs = []
        sub_used = set()
        for sub in item.get('subsections') or []:
            if not isinstance(sub, dict):
                continue
            sub_label = str(sub.get('label') or '').strip()
            if not sub_label:
                continue
            subs.append({
                'key': _unique_key(str(sub.get('key') or _slug(sub_label)), sub_used),
                'label': sub_label,
                'required': bool(sub.get('required')),
                'table_enabled': bool(sub.get('table_enabled')),
                'table_columns': list(sub.get('table_columns') or []),
                'table_column_groups': list(sub.get('table_column_groups') or []),
                'title_region_id': str(sub.get('title_region_id') or ''),
                'content_region_id': str(sub.get('content_region_id') or ''),
                'working_content_region_id': str(sub.get('working_content_region_id') or ''),
            })
        sections.append({
            'key': key,
            'label': label,
            'enabled': item.get('enabled', True) is not False,
            'is_static': bool(item.get('is_static')),
            'static_content': str(item.get('static_content') or ''),
            'table_enabled': bool(item.get('table_enabled')),
            'table_columns': list(item.get('table_columns') or []),
            'table_column_groups': list(item.get('table_column_groups') or []),
            'editor': item.get('editor') or None,
            'title_region_id': str(item.get('title_region_id') or ''),
            'content_region_id': str(item.get('content_region_id') or ''),
            'working_content_region_id': str(item.get('working_content_region_id') or ''),
            'subsections': subs,
        })
    click_log = field_map.get('click_log') if isinstance(field_map.get('click_log'), list) else []
    return {
        'sections': sections,
        'version': 1,
        'click_log': [
            {
                'id': str(entry.get('id') or ''),
                'text': str(entry.get('text') or '')[:500],
                'role': str(entry.get('role') or ''),
                'at': str(entry.get('at') or ''),
            }
            for entry in click_log[-200:]
            if isinstance(entry, dict) and entry.get('id')
        ],
    }
