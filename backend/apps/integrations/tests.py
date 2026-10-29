import json
import os
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase
from rest_framework.test import APIRequestFactory, force_authenticate

from apps.integrations import indexing, llm_service, pii_masking, rag_service, syllabus_structure, syllabus_style
from apps.integrations.syllabus_pipeline import _parse_refine_flag
from apps.integrations.traffic import concurrent_snapshot, record_presence
from apps.integrations.cot_prompt_builder import (
    CoTPromptBuilder,
    assign_topic_weeks,
    complete_syllabus_field,
    field_reply_is_usable,
)
from apps.integrations.llm_service import (
    DEFAULT_CHAIN,
    LLMService,
    message_text,
    parse_reset_delay,
    quota_from_headers,
)
from apps.integrations.models import GlobalDocument
from apps.integrations.views import (
    list_audit_logs,
    list_rag_documents,
    upload_course_rag_document,
    upload_global_rag_document,
)


def _service(*configured):
    """LLMService with only the given providers configured, independent of .env."""
    service = LLMService()
    service.configured = {name: name in configured for name in service.configured}
    service.chain = list(DEFAULT_CHAIN)
    return service


class LLMProviderFallbackTests(SimpleTestCase):
    def setUp(self):
        cache.clear()

    def test_returns_generated_text_by_default(self):
        service = _service('groq')
        service.default_provider = 'groq'

        with patch.object(service, 'generate_response', return_value='Draft text'):
            generated_text = service.generate_response_with_fallback(
                messages=[{'role': 'user', 'content': 'Write a draft'}],
            )

        self.assertEqual(generated_text, 'Draft text')

    def test_walks_the_chain_until_a_provider_succeeds(self):
        service = _service('openrouter', 'groq', 'gemma')
        service.default_provider = 'openrouter'

        with patch.object(
            service,
            'generate_response',
            side_effect=[ValueError('OpenRouter insufficient credits'), ValueError('rate limited'), 'Draft text'],
        ) as generate_response:
            generated_text, provider_used = service.generate_response_with_fallback(
                messages=[{'role': 'user', 'content': 'Write a draft'}],
                max_tokens=1500,
                return_provider=True,
            )

        self.assertEqual(generated_text, 'Draft text')
        self.assertEqual(provider_used, 'gemma')
        self.assertEqual(
            [call.kwargs['provider_choice'] for call in generate_response.call_args_list],
            ['openrouter', 'groq', 'gemma'],
        )

    def test_legacy_and_auto_provider_names(self):
        service = _service('openrouter', 'groq')

        for legacy in ('gemini', 'gpt', 'anthropic'):
            self.assertEqual(
                service.get_provider_and_model(legacy),
                ('openrouter', service.openrouter_model),
            )
        self.assertEqual(service.get_provider_and_model('groq'), ('groq', service.groq_model))
        service.default_provider = 'groq'
        self.assertEqual(service.get_provider_and_model('auto'), ('groq', service.groq_model))

    def test_empty_content_uses_reasoning_text(self):
        self.assertEqual(
            message_text({'content': '', 'reasoning': '{"question_text": "What is RAM?"}'}),
            '{"question_text": "What is RAM?"}',
        )
        self.assertEqual(
            message_text({'content': 'visible answer', 'reasoning': 'hidden reasoning'}),
            'visible answer',
        )

    def test_legacy_alias_does_not_retry_openrouter_twice(self):
        service = _service('openrouter')

        with patch.object(
            service,
            'generate_response',
            side_effect=ValueError('OpenRouter insufficient credits'),
        ) as generate_response:
            with self.assertRaisesRegex(ValueError, 'all available providers'):
                service.generate_response_with_fallback(
                    messages=[{'role': 'user', 'content': 'Write a draft'}],
                    preferred_provider='gemini',
                    max_tokens=1500,
                )

        self.assertEqual(generate_response.call_count, 1)

    def test_rate_limited_provider_is_paused_for_later_requests(self):
        service = _service('groq', 'gemma')
        service.api_keys = {name: 'key' for name in service.api_keys}
        rate_limited = MagicMock(ok=False, status_code=429, text='Rate limit reached', headers={'retry-after': '30'})
        success = MagicMock(ok=True, status_code=200, headers={})
        success.json.return_value = {'choices': [{'message': {'content': 'Draft text'}}]}

        with patch('apps.integrations.llm_service.requests.post', side_effect=[rate_limited, success, success]) as post:
            first = service.generate_response_with_fallback([{'role': 'user', 'content': 'Hi'}], return_provider=True)
            second = service.generate_response_with_fallback([{'role': 'user', 'content': 'Hi'}], return_provider=True)

        self.assertEqual(first, ('Draft text', 'gemma'))
        self.assertEqual(second, ('Draft text', 'gemma'))
        self.assertEqual(post.call_count, 3)
        self.assertGreater(llm_service.provider_cooldown_remaining('groq'), 0)

    def test_large_prompt_skips_groq_and_caps_its_completion_budget(self):
        service = _service('groq', 'gemma')
        big_prompt = 'x' * 4 * 9000

        with patch.object(service, 'generate_response', return_value='Draft text') as generate_response:
            _, provider_used = service.generate_response_with_fallback(
                [{'role': 'user', 'content': big_prompt}], max_tokens=4096, return_provider=True,
            )

        self.assertEqual(provider_used, 'gemma')
        self.assertEqual(generate_response.call_count, 1)
        self.assertEqual(service._completion_budget('groq', 5000, 8192), 2800)

    def test_ollama_takes_over_when_no_online_provider_can_fit_the_prompt(self):
        service = _service('groq', 'gemma')
        service.ollama_configured = True
        big_prompt = 'x' * 4 * 9000

        def no_online_room(provider, prompt_tokens, max_tokens):
            if provider in ('groq', 'gemma', 'cloudflare', 'openrouter'):
                return None
            return max_tokens

        with patch.object(service, '_completion_budget', side_effect=no_online_room), \
             patch.object(service, 'is_ollama_available', return_value=True), \
             patch.object(service, 'generate_response', return_value='Draft text') as generate_response:
            text, provider_used = service.generate_response_with_fallback(
                [{'role': 'user', 'content': big_prompt}], max_tokens=4096, return_provider=True,
            )

        self.assertEqual((text, provider_used), ('Draft text', 'ollama'))
        self.assertEqual(generate_response.call_args.kwargs['provider_choice'], 'ollama')
        self.assertEqual(generate_response.call_count, 1)


class LlmQuotaHealthTests(SimpleTestCase):
    def test_groq_duration_and_unused_tokens(self):
        self.assertAlmostEqual(parse_reset_delay('667ms'), 0.667, places=3)
        self.assertAlmostEqual(parse_reset_delay('1h7m40.8s'), 3600 + 7 * 60 + 40.8, places=1)
        quota = quota_from_headers('groq', {
            'x-ratelimit-limit-tokens': '8000',
            'x-ratelimit-remaining-tokens': '7911',
            'x-ratelimit-limit-requests': '1000',
            'x-ratelimit-remaining-requests': '953',
            'x-ratelimit-reset-tokens': '667ms',
            'x-ratelimit-reset-requests': '1h7m40.8s',
        })
        self.assertEqual(quota['status'], 'healthy')
        self.assertEqual(quota['tokens_remaining'], 7911)
        self.assertEqual(quota['tokens_limit'], 8000)
        self.assertEqual(quota['quota_kind'], 'tokens')

    def test_zero_remaining_tokens_are_a_reached_limit(self):
        quota = quota_from_headers('groq', {
            'x-ratelimit-limit-tokens': '8000',
            'x-ratelimit-remaining-tokens': '0',
            'x-ratelimit-reset-tokens': '30s',
        })
        self.assertEqual(quota['status'], 'limited')
        self.assertIsNotNone(quota['resets_at'])
        self.assertIn('The limit is reached.', quota['detail'])


class ConcurrentTrafficTests(SimpleTestCase):
    def setUp(self):
        cache.clear()

    def test_counts_unique_users_in_the_window(self):
        faculty = SimpleNamespace(user_id='faculty-1', role='faculty')
        admin = SimpleNamespace(user_id='admin-1', role='admin')
        record_presence(faculty)
        record_presence(faculty)
        record_presence(admin)
        snapshot = concurrent_snapshot()
        self.assertEqual(snapshot['concurrent_users'], 2)
        counts = {item['label']: item['count'] for item in snapshot['by_role']}
        self.assertEqual(counts, {'Faculty': 1, 'Admin': 1})
        self.assertEqual(snapshot['window_seconds'], 120)


class SyllabusFieldPromptTests(SimpleTestCase):
    def test_includes_extra_instructions_in_generated_field_prompt(self):
        _, user_prompt = CoTPromptBuilder.build_syllabus_field_prompt(
            course_name='Computer Programming',
            course_code='CS101',
            field_label='Course Outcomes table',
            section='custom',
            existing_content='',
            mode='generate',
            context_chunks=[],
            extra_instructions='Return five rows as a JSON array.',
        )

        self.assertIn('Return five rows as a JSON array.', user_prompt)
        self.assertNotIn('Be concise but complete', user_prompt)

    def test_syllabus_prompt_asks_only_for_course_cells(self):
        system, user = CoTPromptBuilder.build_syllabus_prompt(
            course_name='Science, Technology and Society',
            course_code='GEd 109',
            num_weeks=18,
            context_chunks=[],
            course_category='ged',
            assessment_criteria='25% Midterm Examination\n25% Final Examination',
        )

        self.assertIn('Write only the cells listed, at the length given', system)
        self.assertNotIn('thorough, well-structured syllabi', system)
        self.assertIn('Course category: ged', user)
        self.assertIn('25% Midterm Examination', user)
        self.assertIn('2–3 sentences', user)
        self.assertIn('Three lines', user)
        self.assertIn('week 1 through week 18', user)
        self.assertIn('5-6', user)
        self.assertIn('No tools', user)
        self.assertIn('Do not write Key Learning Activities or Delivery', user)
        for banned in (
            'TABLE OF SPECIFICATIONS',
            'FRAMEWORK ALIGNMENT',
            'COURSE DESCRIPTION (3–4 sentences)',
            'Step 6',
            'Step 7',
            'Step 8',
            'Vision',
        ):
            self.assertNotIn(banned, user)
        self.assertIn('Do not print Step 1', user)
        self.assertEqual(CoTPromptBuilder.SYLLABUS_MAX_TOKENS, 900)

    def test_capstone_stays_professional_and_swaps_the_block(self):
        _, user = CoTPromptBuilder.build_syllabus_prompt(
            course_name='Capstone Project 1',
            course_code='IT 324',
            num_weeks=18,
            context_chunks=[],
            course_category='professional',
        )
        self.assertIn('Course category: professional', user)
        self.assertIn('proposal, completed system, documentation, defense', user)
        self.assertNotIn('Do not emit 18 topic rows', user)

    def test_internship_does_not_ask_for_a_week_per_row(self):
        _, user = CoTPromptBuilder.build_syllabus_prompt(
            course_name='Internship Training',
            course_code='IT 421',
            num_weeks=18,
            context_chunks=[],
            course_category='professional',
        )
        self.assertIn('Course category: professional', user)
        self.assertIn('Do not emit 18 topic rows', user)
        self.assertNotIn('Weekly blocks to emit: 18', user)

    def test_field_hints_match_the_length_rules(self):
        _, tla = CoTPromptBuilder.build_syllabus_field_prompt(
            course_name='Information Management',
            course_code='IT 221',
            field_label='Teaching, Learning, and Assessment Strategies',
            section='tlas',
            existing_content='',
            mode='generate',
            context_chunks=[],
            course_category='common',
            assessment_criteria='40% Laboratory Activities\n30% Case Study',
        )
        self.assertIn('Using only these criteria: 40% Laboratory Activities', tla)
        self.assertIn('Do not add a component', tla)
        self.assertNotIn('comprehensive', tla.lower())

        _, ilo = CoTPromptBuilder.build_syllabus_field_prompt(
            course_name='Information Management',
            course_code='IT 221',
            field_label='Intended Learning Outcome 1',
            section='ilo',
            existing_content='',
            mode='generate',
            context_chunks=[],
            course_category='common',
        )
        self.assertIn('Write one intended learning outcome for this course.', ilo)
        self.assertIn('Do not write "students should be able to"', ilo)
        self.assertIn('Do not mention these rules.', ilo)
        self.assertIn('Define, Apply, Implement, Analyze', ilo)
        self.assertNotIn('TASK:', ilo)
        self.assertEqual(CoTPromptBuilder.field_max_tokens('ilo'), 50)
        self.assertEqual(CoTPromptBuilder.field_max_tokens('rationale'), 160)
        self.assertEqual(CoTPromptBuilder.field_max_tokens('tlas'), 220)
        self.assertEqual(CoTPromptBuilder.field_max_tokens('topics'), 180)
        self.assertEqual(CoTPromptBuilder.field_max_tokens('topic_outcome'), 50)
        self.assertEqual(CoTPromptBuilder.field_max_tokens('topic_map'), 200)
        self.assertEqual(CoTPromptBuilder.field_max_tokens('framework_columns'), 120)
        self.assertIn('framework_columns', CoTPromptBuilder.GENERATABLE_SYLLABUS_SECTIONS)

        _, outcome = CoTPromptBuilder.build_syllabus_field_prompt(
            course_name='Information Management',
            course_code='IT 221',
            field_label='Topic Outcomes 3',
            section='topic_outcome',
            existing_content='',
            mode='generate',
            context_chunks=[],
            course_category='common',
            extra_instructions='Topic: Relational schema.',
        )
        self.assertIn('Write what the student does with this topic.', outcome)
        self.assertIn('Do not mention these rules.', outcome)
        self.assertIn('Topic: Relational schema.', outcome)
        self.assertNotIn('TASK:', outcome)
        self.assertNotIn('Topic Outcomes 3', outcome)

        _, mapped = CoTPromptBuilder.build_syllabus_field_prompt(
            course_name='Information Management',
            course_code='IT 221',
            field_label='ILO and SO mapping',
            section='topic_map',
            existing_content='',
            mode='generate',
            context_chunks=[],
            extra_instructions='Intended learning outcomes:\n1. Define a relational schema.',
        )
        self.assertIn('ILO number, tab, SO number', mapped)
        self.assertIn('1. Define a relational schema.', mapped)
        self.assertNotIn('TASK:', mapped)

        _, columns = CoTPromptBuilder.build_syllabus_field_prompt(
            course_name='Information Management',
            course_code='IT 221',
            field_label='Framework column numbers',
            section='framework_columns',
            existing_content='',
            mode='generate',
            context_chunks=[],
            extra_instructions='SO 5: Design a system.\nILO 1: Design a database.',
        )
        self.assertIn('Do not write the statements again.', columns)
        self.assertIn('SO 5: Design a system.', columns)
        self.assertIn('SO: n, n', columns)
        self.assertTrue(field_reply_is_usable('framework_columns', 'SO: 5\nIGA: none\nCDIO: 2\nSDG: none'))
        self.assertFalse(field_reply_is_usable('framework_columns', 'SO: none\nIGA: none'))

        _, topics = CoTPromptBuilder.build_syllabus_field_prompt(
            course_name='Analytics Techniques and Tools',
            course_code='BAT 404',
            field_label='Topics / Reading List 1',
            section='topics',
            existing_content='',
            mode='generate',
            context_chunks=[],
            course_category='professional',
        )
        self.assertIn('First line is the topic title', topics)
        self.assertIn('Titles only', topics)

    def test_stored_sections_are_not_generated(self):
        for section in ('contact', 'policies', 'attributes', 'assessment_map'):
            self.assertNotIn(section, CoTPromptBuilder.GENERATABLE_SYLLABUS_SECTIONS)
            _, user = CoTPromptBuilder.build_syllabus_field_prompt(
                course_name='Ethics',
                course_code='GEd 107',
                field_label=section,
                section=section,
                existing_content='',
                mode='generate',
                context_chunks=[],
            )
            self.assertIn('Do not generate', user)


class TopicWeekPlanTests(SimpleTestCase):
    def test_single_weeks_run_through_week_18(self):
        raw = "\n".join(f"1 Topic {number}" for number in range(1, 19))
        lines = assign_topic_weeks(raw, term_weeks=18).splitlines()
        self.assertEqual(len(lines), 19)
        self.assertEqual(lines[0], "1\tClass Orientation")
        self.assertTrue(lines[1].startswith("1\tTopic 1"))
        self.assertTrue(lines[10].startswith("10\t"))
        self.assertTrue(lines[14].startswith("14\t"))
        self.assertTrue(lines[17].startswith("17\t"))
        self.assertEqual(lines[-1], "18\tFinal Examination and Submission of Requirements")
        self.assertNotIn("Topic 18", assign_topic_weeks(raw, term_weeks=18))

    def test_model_orientation_is_not_counted_as_a_content_week(self):
        raw = "1 Class Orientation\n" + "\n".join(f"1 Topic {number}" for number in range(1, 19))
        lines = assign_topic_weeks(raw, term_weeks=18).splitlines()
        self.assertEqual(lines[0], "1\tClass Orientation")
        self.assertTrue(lines[1].startswith("1\tTopic 1"))
        self.assertEqual(lines[-1], "18\tFinal Examination and Submission of Requirements")
        self.assertEqual(len(lines), 19)

    def test_heavier_topic_is_one_merged_row(self):
        raw = "1 Opening\n2 Customer Analytics\n" + "\n".join(["1 Later"] * 15)
        lines = assign_topic_weeks(raw, term_weeks=18).splitlines()
        labels = [line.split("\t", 1)[0] for line in lines]
        self.assertEqual(labels[0], "1")
        self.assertEqual(lines[0], "1\tClass Orientation")
        self.assertEqual(labels[1], "1")
        self.assertEqual(labels[2], "2-3")
        self.assertIn("Customer Analytics", lines[2])
        self.assertEqual(lines[-1], "18\tFinal Examination and Submission of Requirements")
        self.assertEqual(len(lines), 18)

    def test_short_plan_widens_until_the_term_is_covered(self):
        text = assign_topic_weeks("1 Alpha\n2 Beta\n1 Gamma", term_weeks=6)
        labels = [line.split("\t", 1)[0] for line in text.splitlines()]
        self.assertEqual(labels, ["1", "1-2", "3-4", "5-6"])
        self.assertTrue(text.splitlines()[0].endswith("Class Orientation"))

    def test_internship_plan_is_not_stretched_to_18_weeks(self):
        text = assign_topic_weeks(
            "2 Deployment\n1 Narrative Report",
            term_weeks=18,
            pad_to_term=False,
        )
        labels = [line.split("\t", 1)[0] for line in text.splitlines()]
        self.assertEqual(labels, ["1", "1-2", "3"])
        self.assertIn("Class Orientation", text.splitlines()[0])
        self.assertIn("Deployment", text.splitlines()[1])


class SyllabusFieldTakeoverTests(SimpleTestCase):
    def _messages(self):
        return {
            "section": "topic_outcome",
            "messages": [{"role": "user", "content": "Topic: Relational schema."}],
            "system": "Reply with the sentence only.",
        }

    def test_a_usable_reply_is_kept(self):
        llm = MagicMock()
        llm.generate_response_with_fallback.return_value = (
            "Design a relational schema for a sample database.",
            "groq",
        )
        text, provider = complete_syllabus_field(llm, **self._messages())
        self.assertEqual(text, "Design a relational schema for a sample database.")
        self.assertEqual(provider, "groq")
        self.assertEqual(llm.generate_response_with_fallback.call_count, 1)
        self.assertFalse(field_reply_is_usable(
            "topic_outcome",
            "We need one sentence, imperative verb, 7-12 words.",
        ))

    def test_unusable_reply_retries_once_then_uses_ollama(self):
        llm = MagicMock()
        echo = "We need one sentence, imperative verb, 7-12 words."
        llm.generate_response_with_fallback.side_effect = [
            (echo, "groq"),
            (echo, "groq"),
            ("Design a relational schema for a sample database.", "ollama"),
        ]
        text, provider = complete_syllabus_field(llm, **self._messages())
        self.assertEqual(text, "Design a relational schema for a sample database.")
        self.assertEqual(provider, "ollama")
        self.assertEqual(
            [call.kwargs["preferred_provider"] for call in llm.generate_response_with_fallback.call_args_list],
            [None, None, "ollama"],
        )

    def test_a_provider_error_goes_straight_to_ollama(self):
        llm = MagicMock()
        llm.generate_response_with_fallback.side_effect = [
            ValueError("all available providers"),
            ("Design a relational schema for a sample database.", "ollama"),
        ]
        text, provider = complete_syllabus_field(llm, **self._messages())
        self.assertEqual(provider, "ollama")
        self.assertIn("relational schema", text)
        self.assertEqual(llm.generate_response_with_fallback.call_count, 2)
        self.assertEqual(llm.generate_response_with_fallback.call_args_list[1].kwargs["preferred_provider"], "ollama")


class InstitutionScopedAdminDataTests(SimpleTestCase):
    def setUp(self):
        self.factory = APIRequestFactory()
        self.admin = SimpleNamespace(
            is_authenticated=True,
            institution_id='institution-a',
        )

    @patch('apps.integrations.views._is_admin_user', return_value=True)
    @patch('apps.integrations.views.GlobalDocument.objects.filter')
    def test_knowledge_base_list_only_queries_admin_institution(
        self,
        filter_documents,
        _is_admin,
    ):
        filter_documents.return_value.order_by.return_value = []
        request = self.factory.get('/api/integrations/rag-documents/')
        force_authenticate(request, user=self.admin)

        response = list_rag_documents(request)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['documents'], [])
        filter_documents.assert_called_once_with(institution_id='institution-a')
        self.assertTrue(_is_admin.called)

    @patch('apps.integrations.views._is_admin_user', return_value=True)
    @patch('apps.integrations.views.AuditLog.objects.select_related')
    def test_audit_log_list_only_queries_admin_institution(
        self,
        select_related,
        _is_admin,
    ):
        select_related.return_value.filter.return_value.order_by.return_value = []
        request = self.factory.get('/api/integrations/audit-logs/')
        force_authenticate(request, user=self.admin)

        response = list_audit_logs(request)

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data['logs'], [])
        select_related.assert_called_once_with('user')
        select_related.return_value.filter.assert_called_once_with(
            user__institution_id='institution-a',
        )
        self.assertTrue(_is_admin.called)


class BackgroundIndexingTests(SimpleTestCase):
    @patch('apps.integrations.views.transaction.on_commit', side_effect=lambda callback: callback())
    @patch('apps.integrations.views.indexing.enqueue_global_document')
    @patch('apps.integrations.views.GlobalDocument.objects.create')
    @patch('apps.integrations.views._is_admin_user', return_value=True)
    def test_upload_returns_202_and_queues_indexing(self, _is_admin, create_document, enqueue, _on_commit):
        create_document.return_value = SimpleNamespace(
            document_id='doc-1', status=GlobalDocument.IndexStatus.PROCESSING,
        )
        user = SimpleNamespace(
            is_authenticated=True, institution_id='institution-a', institution='inst',
        )
        upload = SimpleUploadedFile('syllabus.pdf', b'%PDF-1.4 test', content_type='application/pdf')
        request = APIRequestFactory().post(
            '/api/integrations/upload-global-rag-document/',
            {'file': upload, 'file_type': 'pdf', 'doc_type': 'Syllabus'},
            format='multipart',
        )
        force_authenticate(request, user=user)

        response = upload_global_rag_document(request)

        self.assertEqual(response.status_code, 202, response.data)
        self.assertEqual(response.data['status'], 'PROCESSING')
        self.assertEqual(
            create_document.call_args.kwargs['status'], GlobalDocument.IndexStatus.PROCESSING,
        )
        enqueue.assert_called_once_with('doc-1', b'%PDF-1.4 test', 'pdf', 'syllabus.pdf')

    @patch('apps.integrations.views.transaction.on_commit', side_effect=lambda callback: callback())
    @patch('apps.integrations.views.indexing.enqueue_course_file')
    @patch('apps.integrations.views.CourseFile.objects.create')
    @patch('apps.integrations.views._get_faculty_course_or_403')
    @patch('apps.integrations.views._is_faculty_user', return_value=True)
    def test_course_upload_returns_202_and_queues_indexing(
        self, _is_faculty, get_course, create_file, enqueue, _on_commit,
    ):
        get_course.return_value = (SimpleNamespace(course_id='course-1', course_code='GE4'), None)
        create_file.return_value = SimpleNamespace(records_id='rec-1', rag_status='PROCESSING')
        upload = SimpleUploadedFile('module.pdf', b'%PDF-1.4 test', content_type='application/pdf')
        request = APIRequestFactory().post(
            '/api/integrations/upload-course-rag-document/',
            {
                'file': upload, 'file_type': 'pdf',
                'course_id': '6f1c1f5e-8a8e-4c55-9d43-0d6b6d2f5a11', 'module_tag': 'M1',
            },
            format='multipart',
        )
        force_authenticate(request, user=SimpleNamespace(is_authenticated=True))

        response = upload_course_rag_document(request)

        self.assertEqual(response.status_code, 202, response.data)
        self.assertEqual(response.data['rag_status'], 'PROCESSING')
        enqueue.assert_called_once_with('rec-1', b'%PDF-1.4 test', 'pdf', 'module.pdf', 'M1', None)

    @patch('apps.integrations.indexing.close_old_connections')
    @patch('apps.integrations.indexing.GlobalDocument.objects.filter')
    @patch(
        'apps.integrations.indexing.index_global_document',
        side_effect=ValueError('No text could be extracted from the file.'),
    )
    def test_failed_indexing_marks_document_failed(self, _index, filter_documents, _close):
        indexing._run_index_job('doc-1', b'', 'pdf', 'empty.pdf')

        filter_documents.assert_called_once_with(document_id='doc-1')
        filter_documents.return_value.update.assert_called_once_with(
            status=GlobalDocument.IndexStatus.FAILED,
            error_message='No text could be extracted from the file.',
        )


def _hf_response(status_code, payload=None, text=''):
    response = MagicMock(status_code=status_code, text=text)
    response.json.return_value = payload
    return response


@patch.dict(os.environ, {'EMBEDDING_PROVIDER': 'huggingface', 'HF_TOKEN': 'hf_test'})
class HuggingFaceEmbeddingTests(SimpleTestCase):
    @patch('requests.post')
    def test_batches_and_normalizes(self, post):
        post.side_effect = lambda url, **kw: _hf_response(200, [[3.0, 4.0]] * len(kw['json']['inputs']))

        vectors = rag_service.embed_documents_batch([f'chunk {i}' for i in range(40)])

        self.assertEqual(len(vectors), 40)
        self.assertEqual(post.call_count, 2)
        self.assertAlmostEqual(vectors[0][0], 0.6)
        self.assertAlmostEqual(vectors[0][1], 0.8)
        self.assertIn('BAAI/bge-base-en-v1.5', post.call_args.args[0])
        self.assertEqual(post.call_args.kwargs['headers']['Authorization'], 'Bearer hf_test')

    @patch('apps.integrations.rag_service.time.sleep')
    @patch('requests.post')
    def test_retries_cold_start(self, post, _sleep):
        post.side_effect = [_hf_response(503, text='loading'), _hf_response(200, [[1.0, 0.0]])]

        self.assertEqual(rag_service.embed_query('outcomes'), [1.0, 0.0])
        self.assertEqual(post.call_count, 2)
        self.assertTrue(post.call_args.kwargs['json']['inputs'][0].startswith('Represent this sentence'))

    @patch('requests.post')
    def test_error_is_reported(self, post):
        post.return_value = _hf_response(401, text='Invalid credentials')

        with self.assertRaisesRegex(ValueError, '401'):
            rag_service.embed_documents_batch(['x'])

    def test_missing_token_is_reported(self):
        with patch.dict(os.environ, {'HF_TOKEN': ''}):
            with self.assertRaisesRegex(ValueError, 'HF_TOKEN'):
                rag_service.embed_documents_batch(['x'])


@patch.dict(os.environ, {
    'EMBEDDING_PROVIDER': 'cloudflare,huggingface',
    'CLOUDFLARE_ACCOUNT_ID': 'acct',
    'CLOUDFLARE_API_TOKEN': 'cf_test',
    'HF_TOKEN': 'hf_test',
})
class CloudflareEmbeddingTests(SimpleTestCase):
    @patch('requests.post')
    def test_uses_same_bge_model_on_cloudflare(self, post):
        post.side_effect = lambda url, **kw: _hf_response(
            200, {'result': {'data': [[3.0, 4.0]] * len(kw['json']['text'])}},
        )

        vectors = rag_service.embed_documents_batch([f'chunk {i}' for i in range(150)])

        self.assertEqual(len(vectors), 150)
        self.assertEqual(post.call_count, 2)
        self.assertAlmostEqual(vectors[0][1], 0.8)
        self.assertTrue(post.call_args.args[0].endswith('/accounts/acct/ai/run/@cf/baai/bge-base-en-v1.5'))

    @patch('requests.post')
    def test_falls_back_to_next_provider(self, post):
        post.side_effect = [
            _hf_response(401, text='Authentication error'),
            _hf_response(200, [[1.0, 0.0]]),
        ]

        self.assertEqual(rag_service.embed_documents_batch(['x']), [[1.0, 0.0]])
        self.assertIn('huggingface', post.call_args.args[0])

    @patch('requests.post')
    def test_reports_every_failure(self, post):
        post.side_effect = [
            _hf_response(401, text='Authentication error'),
            _hf_response(402, text='You have no remaining credits'),
        ]

        with self.assertRaisesRegex(ValueError, 'Cloudflare.*401.*Hugging Face.*402'):
            rag_service.embed_documents_batch(['x'])


SAMPLE_SYLLABUS_MD = """
# STATE UNIVERSITY

# Course Specifications

| Course Title | Analytics |
|---|---|
| Course Code | BAT 404 |
| Lab Hours | 3 |
| Consultation Hours | Mon 1-3 PM |

# University Direction

## Vision

A premier national university.

## Mission

A university committed to producing leaders.

# Weekly Schedule

| Chapter | Topics | Weeks | Delivery |
|---|---|---|---|
| 1 | Orientation | 1 | F2F |

# Course Policies

## Attendance

Be on time.

# Laboratory Requirements

| Activity | Tool | Deadline |
|---|---|---|
| LA 1 | Python | Week 3 |
"""


class SyllabusStructureTests(SimpleTestCase):
    def _by_label(self, sections):
        return {section['label']: section for section in sections}

    def test_default_extract_skips_llm(self):
        llm = MagicMock()
        sections, source = syllabus_structure.extract_syllabus_structure(SAMPLE_SYLLABUS_MD, llm=llm)
        self.assertEqual(source, 'outline')
        self.assertTrue(sections)
        llm.generate_response_with_fallback.assert_not_called()

    def test_outline_maps_headings_tables_and_builtin_editors(self):
        sections, source = syllabus_structure.extract_syllabus_structure(SAMPLE_SYLLABUS_MD, refine=False)
        by_label = self._by_label(sections)

        self.assertEqual(source, 'outline')
        self.assertEqual(
            list(by_label),
            ['Course Specifications', 'University Direction', 'Weekly Schedule',
             'Course Policies', 'Laboratory Requirements'],
        )
        specs = by_label['Course Specifications']
        self.assertEqual(specs['key'], 'basic')
        # Fields the built-in Basic Info editor renders are not duplicated.
        self.assertEqual([sub['label'] for sub in specs['subsections']], ['Lab Hours', 'Consultation Hours'])
        self.assertEqual(
            [sub['label'] for sub in by_label['University Direction']['subsections']],
            ['Vision', 'Mission'],
        )
        self.assertEqual(by_label['Weekly Schedule']['key'], 'schedule')
        self.assertFalse(by_label['Weekly Schedule']['table_enabled'])
        self.assertEqual(by_label['Course Policies']['key'], 'policies')
        self.assertEqual(by_label['Course Policies']['subsections'], [])
        lab = by_label['Laboratory Requirements']
        self.assertEqual(lab['key'], 'laboratory-requirements')
        self.assertTrue(lab['table_enabled'])
        self.assertEqual(lab['table_columns'], ['Activity', 'Tool', 'Deadline'])

    def test_refinement_drops_ungrounded_labels_and_copies_static_text(self):
        llm = MagicMock()
        llm.generate_response_with_fallback.return_value = json.dumps({'sections': [
            {'label': 'Course Specifications', 'editor': 'basic', 'is_static': False,
             'table_columns': ['Course Title', 'Course Code', 'Lab Hours', 'Consultation Hours'],
             'subsections': []},
            {'label': 'University Direction', 'editor': None, 'is_static': True,
             'table_columns': [], 'subsections': []},
            {'label': 'Learning Resources Overview', 'editor': None, 'is_static': False,
             'table_columns': [], 'subsections': []},
            {'label': 'Laboratory Requirements', 'editor': None, 'is_static': False,
             'table_columns': ['Activity', 'Tool', 'Invented Column'],
             'subsections': [{'label': 'Imaginary Field', 'table_columns': []}]},
        ]})

        sections, source = syllabus_structure.extract_syllabus_structure(
            SAMPLE_SYLLABUS_MD, refine=True, llm=llm,
        )
        by_label = self._by_label(sections)

        self.assertEqual(source, 'ai')
        self.assertNotIn('Learning Resources Overview', by_label)
        # Label/value labels returned as table columns become subsections.
        specs = by_label['Course Specifications']
        self.assertFalse(specs['table_enabled'])
        self.assertEqual([sub['label'] for sub in specs['subsections']], ['Lab Hours', 'Consultation Hours'])
        direction = by_label['University Direction']
        self.assertTrue(direction['is_static'])
        self.assertIn('A premier national university.', direction['static_content'])
        self.assertIn('Mission:', direction['static_content'])
        lab = by_label['Laboratory Requirements']
        self.assertEqual(lab['table_columns'], ['Activity', 'Tool'])
        self.assertEqual(lab['subsections'], [])

    def test_refinement_failure_falls_back_to_outline(self):
        llm = MagicMock()
        llm.generate_response_with_fallback.side_effect = ValueError('all providers failed')

        sections, source = syllabus_structure.extract_syllabus_structure(
            SAMPLE_SYLLABUS_MD, refine=True, llm=llm,
        )

        self.assertEqual(source, 'outline')
        self.assertEqual(sections[0]['label'], 'Course Specifications')

    def test_refine_prefers_ollama_provider(self):
        llm = MagicMock()
        llm.generate_response_with_fallback.return_value = json.dumps({
            'sections': [{'label': 'Course Specifications', 'editor': 'basic',
                          'is_static': False, 'table_columns': [], 'subsections': []}],
        })
        syllabus_structure.extract_syllabus_structure(
            SAMPLE_SYLLABUS_MD, refine=True, llm=llm, preferred_provider='ollama',
        )
        kwargs = llm.generate_response_with_fallback.call_args.kwargs
        self.assertEqual(kwargs.get('preferred_provider'), 'ollama')


class SyllabusStyleTests(SimpleTestCase):
    def test_docx_style_extracts_margins_and_header_text(self):
        from docx import Document
        from docx.shared import Inches, Pt

        buffer = BytesIO()
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Inches(1)
        section.left_margin = Inches(0.75)
        section.header.paragraphs[0].text = 'University of Example'
        section.footer.paragraphs[0].text = 'Page 1'
        doc.add_paragraph('Body text')
        doc.save(buffer)
        buffer.seek(0)

        profile = syllabus_style.extract_style_profile(buffer, 'docx')
        self.assertEqual(profile['source']['parser'], 'docx')
        self.assertGreaterEqual(profile['source']['confidence'], 0.5)
        self.assertEqual(profile['header']['text'], 'University of Example')
        self.assertIn('Page', profile['footer']['text'])
        self.assertAlmostEqual(profile['page']['margins']['top'], 72.0, delta=1)
        branding = syllabus_style.apply_style_to_branding(profile)
        self.assertEqual(branding['header_text'], 'University of Example')

    def test_parse_refine_flag_defaults_off(self):
        self.assertEqual(_parse_refine_flag(None), (False, None))
        self.assertEqual(_parse_refine_flag('false'), (False, None))
        self.assertEqual(_parse_refine_flag('ollama'), (True, 'ollama'))
        self.assertEqual(_parse_refine_flag('true'), (True, 'ollama'))


class SyllabusDoclingParseTests(SimpleTestCase):
    def _sample_docx(self):
        from docx import Document

        buffer = BytesIO()
        doc = Document()
        doc.add_heading('Course Specifications', level=1)
        doc.add_paragraph('Course Code: CS 101')
        doc.add_paragraph('Lab Hours: 3')
        doc.add_heading('Weekly Schedule', level=1)
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = 'Week'
        table.rows[0].cells[1].text = 'Topic'
        table.rows[1].cells[0].text = '1'
        table.rows[1].cells[1].text = 'Intro'
        doc.save(buffer)
        buffer.seek(0)
        return buffer

    def test_docx_upload_uses_docling_when_installed(self):
        from apps.integrations import syllabus_parse

        self.assertTrue(
            syllabus_parse._docling_available(),
            'Docling must be installed (pip install docling==2.142.1)',
        )
        markdown, parser = syllabus_parse.markdown_from_upload(self._sample_docx(), 'docx')
        self.assertEqual(parser, 'docling')
        self.assertIn('Course Specifications', markdown)
        self.assertIn('Weekly Schedule', markdown)

    def test_docling_tree_builds_outline_and_pipeline_confidence(self):
        from apps.integrations import syllabus_parse, syllabus_docling, syllabus_pipeline

        self.assertTrue(syllabus_parse._docling_available())
        parsed = syllabus_parse.parse_upload(self._sample_docx(), 'docx')
        self.assertEqual(parsed['parser'], 'docling')
        self.assertIsNotNone(parsed['document'])
        outline = syllabus_docling.build_outline_from_docling(parsed['document'])
        labels = [section['label'] for section in outline]
        self.assertIn('Course Specifications', labels)
        self.assertIn('Weekly Schedule', labels)

        result = syllabus_pipeline.analyze_syllabus_template(self._sample_docx(), 'docx', refine='false')
        self.assertEqual(result['parser'], 'docling')
        self.assertEqual(result['structure_source'], 'docling_tree')
        self.assertTrue(result['sections'])
        self.assertTrue(all('confidence' in section for section in result['sections']))


class SyllabusTemplateDocTests(SimpleTestCase):
    def test_preview_click_map_and_fill_roundtrip(self):
        from docx import Document
        from apps.integrations import syllabus_template_doc

        buffer = BytesIO()
        doc = Document()
        doc.add_heading('Vision and Mission', level=1)
        doc.add_paragraph('PLACEHOLDER_VISION')
        doc.add_heading('Course Information', level=1)
        doc.add_paragraph('PLACEHOLDER_COURSE')
        doc.save(buffer)
        template_bytes = buffer.getvalue()

        preview = syllabus_template_doc.extract_docx_preview(BytesIO(template_bytes))
        self.assertGreaterEqual(len(preview['blocks']), 4)
        title = next(b for b in preview['blocks'] if 'Vision' in b['text'])
        content = next(b for b in preview['blocks'] if b['text'] == 'PLACEHOLDER_VISION')

        field_map = {
            'sections': [{
                'key': 'vision-mission',
                'label': 'Vision and Mission',
                'title_region_id': title['id'],
                'content_region_id': content['id'],
                'enabled': True,
                'is_static': False,
                'subsections': [],
            }],
        }
        sections = syllabus_template_doc.field_map_to_syllabus_sections(field_map)
        self.assertEqual(sections[0]['label'], 'Vision and Mission')
        self.assertEqual(sections[0]['content_region_id'], content['id'])

        filled = syllabus_template_doc.fill_docx_from_field_map(
            template_bytes,
            field_map,
            {'vision-mission': 'A premier university.'},
        )
        out = Document(BytesIO(filled))
        texts = [p.text for p in out.paragraphs]
        self.assertIn('A premier university.', texts)

    def test_xlsx_exact_preview_keeps_cells(self):
        import openpyxl
        from apps.integrations import syllabus_template_doc

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Course Specs'
        ws.append(['Field', 'Value'])
        ws.append(['Course Code', 'CS101'])
        xlsx_buf = BytesIO()
        wb.save(xlsx_buf)
        built = syllabus_template_doc.build_exact_preview(BytesIO(xlsx_buf.getvalue()), 'xlsx')
        self.assertEqual(built['kind'], 'xlsx')
        labels = ' '.join(b['text'] for b in built['blocks'])
        self.assertIn('Course Specs', labels)
        self.assertIn('CS101', labels)
        self.assertTrue(built['working_docx_bytes'])

    def test_unsupported_type_rejected(self):
        from apps.integrations import syllabus_template_doc
        with self.assertRaises(ValueError):
            syllabus_template_doc.normalize_upload_to_docx(BytesIO(b'hi'), 'txt')


class SyllabusParseTimeoutTests(SimpleTestCase):
    def test_liteparse_mode_skips_docling(self):
        from docx import Document
        from apps.integrations import syllabus_parse

        buffer = BytesIO()
        doc = Document()
        doc.add_heading('Course Policies', level=1)
        doc.add_paragraph('Be on time.')
        doc.save(buffer)
        buffer.seek(0)

        with patch.dict(os.environ, {'SYLLABUS_PARSER': 'liteparse'}):
            parsed = syllabus_parse.parse_upload(buffer, 'docx')
        self.assertEqual(parsed['parser'], 'liteparse')
        self.assertIsNone(parsed['document'])
        self.assertIn('Course Policies', parsed['markdown'])


class SyllabusExportTests(SimpleTestCase):
    def test_styled_pdf_export_produces_bytes(self):
        from apps.integrations.syllabus_export import build_syllabus_html, export_syllabus_pdf

        syllabus = {
            'course_title': 'Intro to Computing',
            'course_code': 'CS101',
            'university_name': 'Example University',
            'course_description': 'An introductory course.',
        }
        branding = {
            'header_text': 'Example University',
            'footer_text': 'Confidential',
            'syllabus_sections': [],
            'syllabus_style_profile': {
                'page': {'size': 'A4', 'orientation': 'portrait',
                         'margins': {'top': 54, 'right': 54, 'bottom': 54, 'left': 54}},
                'typography': {
                    'body': {'family': 'Times New Roman', 'size_pt': 11},
                    'heading_1': {'family': 'Arial', 'size_pt': 14, 'weight': 'bold'},
                },
                'colors': {'text': '#111111', 'accent': '#003366', 'table_header_bg': '#eeeeee'},
                'table': {'border': True, 'header_bold': True},
            },
        }
        html_doc = build_syllabus_html(syllabus, branding=branding, style_profile=branding['syllabus_style_profile'])
        self.assertIn('Intro to Computing', html_doc)
        self.assertIn('003366', html_doc)
        pdf_bytes, engine = export_syllabus_pdf(syllabus, branding=branding)
        self.assertIn(engine, {'weasyprint', 'reportlab'})
        self.assertTrue(pdf_bytes.startswith(b'%PDF'))


class PiiMaskingTests(SimpleTestCase):
    def test_masks_email_and_ph_mobile(self):
        raw = (
            "Contact the instructor at juan.cruz@batstate-u.edu.ph "
            "or mobile 09171234567 / +63 917 123 4567."
        )
        result = pii_masking.mask_pii(raw)
        self.assertNotIn('juan.cruz@batstate-u.edu.ph', result.text)
        self.assertNotIn('09171234567', result.text)
        self.assertNotIn('917 123 4567', result.text)
        self.assertIn(pii_masking.EMAIL_PLACEHOLDER, result.text)
        self.assertIn(pii_masking.PHONE_PLACEHOLDER, result.text)
        pii_masking.assert_no_high_confidence_pii(result.text)

    def test_masks_signature_block_but_keeps_course_body(self):
        raw = (
            "## Course Description\n"
            "This course covers Newton's laws of motion and Rizal's novels.\n\n"
            "Prepared by:\n"
            "Prof. Juan Dela Cruz\n"
            "Faculty, College of Engineering\n"
            "_______________________________\n"
        )
        result = pii_masking.mask_and_verify(raw, context='test')
        self.assertIn("Newton's laws of motion", result)
        self.assertIn("Rizal's novels", result)
        self.assertNotIn('Juan Dela Cruz', result)
        self.assertIn(pii_masking.SIGNATURE_PLACEHOLDER, result)

    def test_masks_labeled_instructor_name_only(self):
        raw = (
            "Instructor: Maria Santos\n"
            "Week 3 Topic: Darwin and natural selection\n"
            "Reference: works of Einstein on relativity\n"
        )
        result = pii_masking.mask_pii(raw, mask_names=True)
        self.assertIn(pii_masking.NAME_PLACEHOLDER, result.text)
        self.assertNotIn('Maria Santos', result.text)
        self.assertIn('Darwin and natural selection', result.text)
        self.assertIn('Einstein on relativity', result.text)

    def test_pre_embed_check_rejects_unmasked_email(self):
        with self.assertRaises(ValueError) as ctx:
            pii_masking.assert_no_high_confidence_pii(
                'Still has secret@example.com in text',
                context='chunk',
            )
        self.assertIn('email', str(ctx.exception))

    def test_chunk_and_embed_scans_before_embedding(self):
        parents = {
            'parents': [{
                'content': 'ILO: apply Ohm\'s law. Contact x@y.com',
                'children': [{'content': 'Call 09181234567 for consultation.'}],
            }],
        }
        with patch.object(
            indexing.DocumentProcessor, 'process_file_hierarchical', return_value=parents,
        ), patch.object(
            indexing.rag_service, 'embed_documents_batch', return_value=[[0.1, 0.2]],
        ) as embed:
            _parents, child_records, _ = indexing._chunk_and_embed(b'x', 'pdf', 'demo.pdf')

        embedded_text = child_records[0][1]['content']
        self.assertNotIn('09181234567', embedded_text)
        self.assertIn(pii_masking.PHONE_PLACEHOLDER, embedded_text)
        embed.assert_called_once()
        self.assertNotIn('09181234567', embed.call_args.args[0][0])
