"""Fixed-format Course Information Syllabus (CIS) export.

Every faculty syllabus download - DOCX or PDF - is rendered from the same
document model below. The layout matches the university's standard CIS form:
letterhead, course information grid, ILOs, assessment distribution map,
references, IGA/SO/CDIO/SDG statements, course policies, TLA activities,
assessment schedule, mapping of assessment tasks, signatures, and remarks.

Sections are always emitted in that order. Missing inputs become blank cells
instead of disappearing, so an unfinished syllabus still reads as the form.
"""

from __future__ import annotations

import base64
import binascii
import html
import re
from dataclasses import dataclass, field
from io import BytesIO
from xml.sax.saxutils import escape as xml_escape

MAROON = '8B0000'
HEADER_FILL = 'F2F2F2'

AT_NOTE = (
    'All internal assessments with feedback will be made available within 2 week '
    'after each assessment submission except Final Project.'
)
GRADING_FOOTNOTE = (
    '*Students who got a computed grade of 70-74 will be given an appropriate remedial '
    'activity in which the final grade should be either passing (3.0) or failure (5.0). '
    'No rounding off of grades shall be allowed.'
)
REMARKS = [
    'The syllabus is to be distributed to the students in the first week of the semester.',
    'Any changes to the syllabus shall be communicated (in writing) to the Program Chair '
    'and the approved revised version must be communicated to the students.',
    "The course instructor may set a more stringent similarity percentage (minimum 20%) "
    "for their respective courses pertaining to student's submissions. However, it must be "
    'communicated in writing to the respective Program Chair and the approved revised version '
    'must be communicated to the students.',
]
# Fixed CIS grading bands. The labels always print; grade/range stay blank until
# the faculty fills them in.
GRADING_BANDS = [
    'Excellent', 'Superior', 'Very Good', 'Good', 'Meritorious',
    'Very Satisfactory', 'Satisfactory', 'Fairly Satisfactory',
    'Passing', 'Failure', 'Incomplete',
]
# The three header lines are fixed for every export. Course records sometimes
# carry campus/address text in these fields, which must not reach the document.
FIXED_COLLEGE = 'College of Informatics and Computing Sciences'
FIXED_PROGRAM = 'Bachelor of Science in Information Technology'
FIXED_TITLE = 'COURSE INFORMATION SYLLABUS (CIS)'


# ── document model ──────────────────────────────────────────────────────────


@dataclass
class Cell:
    text: str = ''
    colspan: int = 1
    rowspan: int = 1
    bold: bool = False
    italic: bool = False
    align: str = 'left'
    shade: bool = False
    first_line_bold: bool = False
    rest_italic: bool = False


@dataclass
class Table:
    rows: list[list[Cell]]
    widths: list[float]
    header_rows: int = 0
    size: float = 8.0
    no_split: bool = False
    keep_together: bool = False


@dataclass
class Paragraph:
    text: str = ''
    bold: bool = False
    italic: bool = False
    align: str = 'left'
    size: float = 9.0


@dataclass
class Letterhead:
    image: str = ''
    lines: list[tuple[str, dict]] = field(default_factory=list)


@dataclass
class StandardDocument:
    letterhead: Letterhead
    blocks: list
    college: str = 'College of Informatics and Computing Sciences'
    program: str = 'Bachelor of Science in Information Technology'
    title: str = 'COURSE INFORMATION SYLLABUS (CIS)'


# ── value helpers ───────────────────────────────────────────────────────────


def _s(value):
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'Yes' if value else 'No'
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _first(data, *keys):
    for key in keys:
        value = _s(data.get(key))
        if value:
            return value
    return ''


def _as_lines(value):
    if isinstance(value, (list, tuple)):
        return [_s(item) for item in value if _s(item)]
    text = _s(value)
    return [line.strip() for line in text.splitlines() if line.strip()]


def _pad(values, size):
    values = list(values or [])
    return (values + [''] * size)[:size]


def _numeric_sum(values):
    total = 0.0
    found = False
    for value in values:
        text = _s(value).replace('%', '').replace(',', '')
        try:
            total += float(text)
            found = True
        except (TypeError, ValueError):
            continue
    if not found:
        return ''
    return str(int(total)) if total.is_integer() else f'{total:g}'


def _course_category(value):
    token = _s(value).strip().lower()
    return {
        'ged': 'General Education',
        'general education': 'General Education',
        'common': 'Common Course',
        'common courses': 'Common Course',
        'professional': 'Professional Course',
        'professional courses': 'Professional Course',
        'mandated': 'Mandated Course',
        'mandated courses': 'Mandated Course',
    }.get(token, _s(value))


def _contact_hours(data):
    direct = _s(data.get('contact_hours'))
    if direct:
        return direct
    parts = []
    lecture = _s(data.get('lecture_hours'))
    lab = _s(data.get('lab_hours') or data.get('laboratory_hours'))
    total = _s(data.get('total_hours'))
    if lecture:
        parts.append(f'{lecture} Hours Lecture')
    if lab:
        parts.append(f'{lab} Hours Laboratory')
    if total and not (lecture or lab):
        parts.append(f'{total} Hours Total')
    return '\n'.join(parts)


def _rationale(data):
    rationale = _s(data.get('rationale'))
    description = _s(data.get('description') or data.get('course_description'))
    if rationale and description and description not in rationale:
        return f'{rationale}\n\n{description}'
    return rationale or description


def _normalize_ilos(data):
    items = data.get('ilos') or []
    ilos = []
    for item in items:
        if isinstance(item, dict):
            ilos.append(_first(item, 'description', 'text', 'ilo', 'statement', 'title'))
        else:
            ilos.append(_s(item))
    return ilos


def _normalize_references(data):
    items = data.get('references') or []
    refs = []
    for item in items:
        if isinstance(item, dict):
            refs.append(_first(item, 'citation', 'text', 'reference', 'title', 'description'))
        else:
            refs.append(_s(item))
    return [ref for ref in refs if ref]


def _normalize_framework(data, key, prefix, count):
    items = data.get(key) or []
    total = max(count, len(items))
    normalized = []
    for index in range(total):
        item = items[index] if index < len(items) else {}
        if isinstance(item, str):
            item = {'description': item}
        code = _first(item, 'code') or f'{prefix}{index + 1}'
        title = _first(item, 'title', 'name', 'label')
        description = _first(item, 'description', 'content', 'text', 'statement')
        normalized.append({'code': code, 'title': title, 'description': description})
    return normalized


def _normalize_tasks(data):
    tasks = []
    for item in data.get('assessment_tasks') or []:
        if not isinstance(item, dict):
            continue
        ilo_values = item.get('ilo_values')
        if not isinstance(ilo_values, list):
            ilo_values = [
                _s(item.get(f'ilo{index}')) for index in range(1, 10)
            ]
            while ilo_values and not ilo_values[-1]:
                ilo_values.pop()
        tasks.append({
            'code': _first(item, 'code'),
            'task': _first(item, 'task', 'assessment_task', 'assessment', 'title', 'label'),
            'ird': _first(item, 'ird', 'i_r_d'),
            'weight': _first(item, 'weight', 'percentage', 'percent'),
            'ilo_values': [_s(value) for value in ilo_values],
            'c': _first(item, 'c'),
            'p': _first(item, 'p'),
            'a': _first(item, 'a'),
        })
    return tasks


def _normalize_chapters(data):
    chapters = []
    for item in data.get('tla_chapters') or []:
        if not isinstance(item, dict):
            continue
        topics = item.get('topics')
        if isinstance(topics, list):
            topics = '\n'.join(_s(value) for value in topics if _s(value))
        chapters.append({
            'chapter': _first(item, 'chapter', 'ch', 'chapter_no'),
            'topics': _s(topics),
            'weeks': _first(item, 'weeks', 'week', 'wks'),
            'outcomes': _first(item, 'outcomes', 'topic_outcomes', 'outcome'),
            'ilo': _first(item, 'ilo', 'ilo_refs'),
            'so': _first(item, 'so'),
            'delivery': _first(item, 'delivery', 'delivery_method'),
        })
    return chapters


def _normalize_schedule(data):
    methods = []
    for item in data.get('schedule_methods') or []:
        if not isinstance(item, dict):
            continue
        methods.append({
            'name': _first(item, 'name', 'method', 'label'),
            'values': _pad([_s(value) for value in (item.get('values') or [])], 16),
        })
    return methods


def _normalize_grading(data):
    provided = []
    for item in data.get('grading_rows') or []:
        if not isinstance(item, dict):
            continue
        provided.append({
            'grade': _first(item, 'grade'),
            'rating': _first(item, 'rating', 'description', 'remark'),
            'range': _first(item, 'range', 'grade_range'),
        })
    if not provided:
        return [{'grade': '', 'rating': label, 'range': ''} for label in GRADING_BANDS]
    rows = []
    for index, label in enumerate(GRADING_BANDS):
        item = provided[index] if index < len(provided) else {}
        rows.append({
            'grade': item.get('grade', ''),
            'rating': item.get('rating') or label,
            'range': item.get('range', ''),
        })
    for item in provided[len(GRADING_BANDS):]:
        rows.append(item)
    return rows


def _normalize_mapping(data, key):
    rows = []
    for item in data.get(key) or []:
        if not isinstance(item, dict):
            continue
        rows.append(item)
    return rows


def _mapping_headers(data, key, count):
    """Configured column numbers, padded so the fixed form keeps its column count."""
    values = [_s(value) for value in (data.get(key) or [])]
    while values and not values[-1]:
        values.pop()
    if len(values) < count:
        values.extend([''] * (count - len(values)))
    return values


# ── cell constructors ───────────────────────────────────────────────────────


def _cell(text='', **kwargs):
    return Cell(text=_s(text), **kwargs)


def _label(text, rowspan=1):
    return Cell(text=_s(text), rowspan=rowspan)


def _head(text, colspan=1, rowspan=1, align='center'):
    return Cell(text=_s(text), colspan=colspan, rowspan=rowspan, bold=True, align=align, shade=True)


def _table_columns(rows):
    occupied = set()
    columns = 0
    for row_index, row in enumerate(rows):
        column = 0
        for cell in row:
            while (row_index, column) in occupied:
                column += 1
            for row_span in range(row_index, row_index + cell.rowspan):
                for col_span in range(column, column + cell.colspan):
                    occupied.add((row_span, col_span))
            column += cell.colspan
            columns = max(columns, column)
    return columns


# ── document assembly ───────────────────────────────────────────────────────


def _letterhead(syllabus, branding):
    image = _s(branding.get('header_image') or branding.get('header_image_url'))
    if image:
        # The admin uploads a single header template per institution. It is the
        # whole letterhead, so no generated text is layered next to it.
        return Letterhead(image=image, lines=[])

    raw = [
        segment.strip()
        for segment in re.split(r'[|\n]+', _s(branding.get('header_text')))
        if segment.strip()
    ]
    lowered = ' | '.join(raw).lower()
    lines = []

    def add(text, **style):
        text = _s(text)
        if not text:
            return
        if any(text.lower() == existing[0].lower() for existing in lines):
            return
        lines.append((text, style))

    add('Republic of the Philippines', size=9)

    university = _s(syllabus.get('university_name'))
    if not university:
        university = next((line for line in raw if 'university' in line.lower()), '')
    university = university or 'BATANGAS STATE UNIVERSITY'
    add(university.upper() if university == university.upper() or len(university) <= 40 else university,
        size=13, bold=True)

    tagline = _s(syllabus.get('university_tagline'))
    if not tagline:
        tagline = next((line for line in raw if 'national engineering university' in line.lower()), '')
    if not tagline and 'batangas' in university.lower():
        tagline = 'The National Engineering University'
    if tagline:
        add(tagline, size=10, bold=True, color=MAROON)

    used = {university.lower(), tagline.lower(), 'republic of the philippines'}
    for line in raw:
        if line.lower() in used:
            continue
        add(line, size=8.5)

    for key in ('campus', 'institution_campus', 'address', 'institution_address',
                'institution_contact', 'telephone', 'contact_number'):
        add(syllabus.get(key), size=8.5)

    email = _first(syllabus, 'institution_email', 'email_address')
    website = _s(syllabus.get('website') or syllabus.get('website_address'))
    if email or website:
        parts = []
        if email:
            parts.append(f'E-mail Address: {email}')
        if website:
            parts.append(f'Website Address: {website}')
        add(' | '.join(parts), size=8.5)

    return Letterhead(image='', lines=lines)


def _course_info_table(data):
    instructor_rows = [
        [
            _label('Course Instructor', rowspan=3),
            _cell(_first(data, 'instructor_name')),
            _cell(_first(data, 'employee_id')),
            _cell('Reference CMO'),
            _cell(_first(data, 'reference_cmo'), colspan=2),
        ],
        [
            _cell(_first(data, 'education', 'instructor_qualification'), colspan=2),
            _cell('Date Prepared'),
            _cell(_first(data, 'date_prepared'), colspan=2),
        ],
        [
            _cell(_first(data, 'instructor_email', 'email'), colspan=2),
            _cell('Revision No.'),
            _cell(_first(data, 'revision_no'), colspan=2),
        ],
    ]
    rows = [
        [_label('Vision'), _cell(_s(data.get('vision')), colspan=5)],
        [_label('Mission'), _cell(_s(data.get('mission')), colspan=5)],
        [
            _label('Course Title'),
            _cell(_first(data, 'course_title'), colspan=2),
            _label('Course Code'),
            _cell(_first(data, 'course_code'), colspan=2),
        ],
        [
            _label('Course Category'),
            _cell(_course_category(data.get('course_category')), colspan=2),
            _label('Pre-requisite(s)'),
            _cell(_first(data, 'prerequisites', 'pre_requisites'), colspan=2),
        ],
        [
            _label('Semester/Year'),
            _cell(_first(data, 'semester_year', 'semester'), colspan=2),
            _label('Credit Hours'),
            _cell(_first(data, 'credit_hours', 'units'), colspan=2),
        ],
        *instructor_rows,
        [
            _label('Period of Study'),
            _cell(_first(data, 'period_of_study'), colspan=2),
            _label('Revision Date'),
            _cell(_first(data, 'revision_date'), colspan=2),
        ],
        [
            _label('Course Rationale and Description'),
            _cell(_rationale(data), colspan=5),
        ],
        [_label('Contact Hours'), _cell(_contact_hours(data), colspan=5)],
        [_label('Criteria for Assessment'), _cell(_first(data, 'assessment_criteria'), colspan=5)],
        [
            _label('Teaching, Learning, and Assessment Strategies'),
            _cell(_first(data, 'tlas', 'teaching_strategies', 'teaching_learning_strategies'), colspan=5),
        ],
    ]
    return Table(rows=rows, widths=[1.15, 1.85, 0.65, 1.0, 0.9, 1.65], size=8, keep_together=True)


def _ilo_table(data):
    ilos = _normalize_ilos(data)
    count = max(3, len(ilos))
    rows = [[
        _label('Intended Learning Outcomes (ILO)', rowspan=count + 1),
        _head('ILO'),
        _head('Upon completion of this course, the students should be able to:', align='left'),
    ]]
    for index in range(count):
        rows.append([
            Cell(text=f'ILO{index + 1}', bold=True, align='center'),
            _cell(ilos[index] if index < len(ilos) else ''),
        ])
    return Table(rows=rows, widths=[1.25, 0.55, 5.4], header_rows=1, size=8.5, keep_together=True)


def _at_table(data):
    tasks = _normalize_tasks(data)
    ilo_count = max(3, len(_normalize_ilos(data)), max((len(task['ilo_values']) for task in tasks), default=0))
    total_columns = 1 + 4 + ilo_count + 3
    rows = [
        [
            _label('Assessment Method and Distribution Map', rowspan=2 + len(tasks) + 2),
            _head('Assessment Tasks (AT) Distribution', colspan=4),
            _head('Intended Learning Outcomes', colspan=ilo_count),
            _head('Domains', colspan=3),
        ],
        [
            _head('Code'),
            _head('Assessment Tasks'),
            _head('I/R/D'),
            _head('(%)'),
            *[_head(str(index + 1)) for index in range(ilo_count)],
            _head('C'),
            _head('P'),
            _head('A'),
        ],
    ]
    for task in tasks:
        rows.append([
            _cell(task['code']),
            _cell(task['task']),
            _cell(task['ird']),
            Cell(text=task['weight'], align='center'),
            *[Cell(text=value, align='center') for value in _pad(task['ilo_values'], ilo_count)],
            Cell(text=task['c'], align='center'),
            Cell(text=task['p'], align='center'),
            Cell(text=task['a'], align='center'),
        ])
    ilo_totals = [
        _numeric_sum([
            task['ilo_values'][index] if index < len(task['ilo_values']) else ''
            for task in tasks
        ])
        for index in range(ilo_count)
    ]
    rows.append([
        Cell(text='Total', colspan=3, bold=True),
        Cell(text=_numeric_sum([task['weight'] for task in tasks]), align='center', bold=True),
        *[Cell(text=value, align='center', bold=True) for value in ilo_totals],
        Cell(), Cell(), Cell(),
    ])
    rows.append([
        Cell(text='Note:', bold=True),
        Cell(text=AT_NOTE, colspan=total_columns - 2),
    ])

    fixed = [1.25, 0.45, 1.95, 0.5, 0.45]
    domain_width = 0.42
    remaining = 7.2 - sum(fixed) - (domain_width * 3)
    ilo_width = remaining / ilo_count if ilo_count else 0.3
    widths = fixed + [ilo_width] * ilo_count + [domain_width] * 3
    return Table(rows=rows, widths=widths, header_rows=2, size=8, keep_together=True)


def _references_table(data):
    references = _normalize_references(data)
    count = max(5, len(references))
    rows = []
    for index in range(count):
        row = [_label('Books and Other References', rowspan=count)] if index == 0 else []
        row.extend([
            Cell(text=str(index + 1), align='center'),
            _cell(references[index] if index < len(references) else ''),
        ])
        rows.append(row)
    return Table(rows=rows, widths=[1.25, 0.4, 5.55], size=8.5, no_split=True, keep_together=True)


def _framework_table(data, key, prefix, count, label, header, statement_header):
    items = _normalize_framework(data, key, prefix, count)
    rows = [[
        _label(label, rowspan=len(items) * 2 + 1),
        _head(header),
        _head(statement_header, align='center'),
    ]]
    for item in items:
        rows.append([
            Cell(text=item['code'], align='center', bold=True, rowspan=2),
            Cell(text=item['title'], bold=True),
        ])
        rows.append([_cell(item['description'])])
    return Table(rows=rows, widths=[1.25, 0.55, 5.4], header_rows=1, size=8, keep_together=True)


def _policies_table(data):
    rows = [
        [Cell(text='COURSE POLICIES', colspan=4, bold=True)],
        [Cell(text='A.', align='center'), Cell(text='GRADING SYSTEM', colspan=3, bold=True)],
        [Cell(), Cell(text='The grading system adopted by this course is as follows:', colspan=3)],
    ]
    grading = _normalize_grading(data)
    if grading:
        for row in grading:
            rows.append([
                Cell(),
                _cell(row['rating']),
                Cell(text=row['grade'], align='center'),
                _cell(row['range']),
            ])
    else:
        rows.append([Cell(), Cell(), Cell(), Cell()])
    rows.extend([
        [Cell(), Cell(text=GRADING_FOOTNOTE, colspan=3)],
        [Cell(text='B.'), Cell(text='CLASS POLICY', colspan=3, bold=True)],
        [Cell(), Cell(text=_first(data, 'class_policy'), colspan=3)],
        [Cell(), Cell(text='MISSED EXAMINATIONS', colspan=3, bold=True)],
        [Cell(), Cell(text=_first(data, 'missed_examinations'), colspan=3)],
        [Cell(), Cell(text='ACADEMIC DISHONESTY', colspan=3, bold=True)],
        [Cell(), Cell(text=_first(data, 'academic_dishonesty'), colspan=3)],
        [Cell(), Cell(text='DROPPING', colspan=3, bold=True)],
        [Cell(), Cell(text=_first(data, 'dropping'), colspan=3)],
        [Cell(text='C.'), Cell(text='OTHER COURSE POLICIES AND REQUIREMENTS', colspan=3, bold=True)],
        [Cell(), Cell(text=_first(data, 'other_policies'), colspan=3)],
    ])
    return Table(rows=rows, widths=[0.35, 1.2, 0.85, 4.8], size=8, keep_together=True)


def _tla_table(data):
    chapters = _normalize_chapters(data) or [{
        'chapter': '', 'topics': '', 'weeks': '', 'outcomes': '',
        'ilo': '', 'so': '', 'delivery': '',
    }]
    rows = [[
        Cell(
            text='Teaching, Learning, and Assessment (TLA) Activities',
            colspan=7, bold=True, align='center', shade=True,
        ),
    ], [
        _head('Ch.'),
        _head('Topics / Reading List'),
        _head('Wks'),
        _head('Topic Outcomes'),
        _head('ILO'),
        _head('SO'),
        _head('Delivery Method'),
    ]]
    for chapter in chapters:
        rows.append([
            Cell(text=chapter['chapter'], align='center', bold=True),
            Cell(text=chapter['topics'], first_line_bold=True, rest_italic=True),
            Cell(text=chapter['weeks'], align='center'),
            _cell(chapter['outcomes']),
            Cell(text=chapter['ilo'], align='center'),
            Cell(text=chapter['so'], align='center'),
            _cell(chapter['delivery']),
        ])
    return Table(
        rows=rows,
        widths=[0.4, 1.95, 0.45, 1.55, 0.4, 0.4, 2.05],
        header_rows=2,
        size=8,
    )


def _schedule_table(data):
    methods = _normalize_schedule(data) or [{'name': '', 'values': [''] * 16}]
    rows = [
        [Cell(text='Assessment Method', rowspan=2), _head('Assessment Schedule', colspan=17)],
        [_head('Distribution'), *[_head(str(week)) for week in range(2, 18)]],
    ]
    for method in methods:
        rows.append(
            [_cell(method['name'])]
            + [Cell(text=value, align='center') for value in _pad(method['values'], 16)]
        )
    widths = [0.45, 0.75] + [0.375] * 16
    return Table(rows=rows, widths=widths, header_rows=2, size=8)


def _so_mapping_table(data):
    headers = _mapping_headers(data, 'so_headers', 5)
    source_rows = _normalize_mapping(data, 'so_cpa_rows')
    title = 'STUDENT OUTCOMES (SO): Mapping of Assessment Tasks (AT)'
    rows = [[
        _label('ILO-SO and ILO-CPA Mapping', rowspan=2),
        _head('ILOs', rowspan=2),
        _head(title, colspan=len(headers) + 3, align='center'),
    ], [
        *[_head(header) for header in headers],
        _head('C'), _head('P'), _head('A'),
    ]]
    for index in range(7):
        source = source_rows[index] if index < len(source_rows) else {}
        values = _pad(source.get('so_vals') or [], len(headers))
        rows.append([
            Cell(),
            Cell(text=f'ILO{index + 1}', align='center', bold=True),
            *[Cell(text=value, align='center') for value in values],
            Cell(text=_s(source.get('c')), align='center'),
            Cell(text=_s(source.get('p')), align='center'),
            Cell(text=_s(source.get('a')), align='center'),
        ])
    data_width = 5.8 / (len(headers) + 3)
    widths = [0.9, 0.5] + [data_width] * (len(headers) + 3)
    return Table(rows=rows, widths=widths, header_rows=2, size=7.5, keep_together=True)


def _iga_mapping_table(data):
    headers = _mapping_headers(data, 'iga_headers', 8)
    source_rows = _normalize_mapping(data, 'iga_rows')
    title = 'INSTITUTIONAL GRADUATE ATTRIBUTES (IGA): Mapping of Assessment Tasks (AT)'
    rows = [[
        _label('ILO-IGA Mapping', rowspan=2),
        _head('ILOs', rowspan=2),
        _head(title, colspan=len(headers), align='center'),
    ], [_head(header) for header in headers]]
    for index in range(7):
        source = source_rows[index] if index < len(source_rows) else {}
        values = _pad(source.get('vals') or [], len(headers))
        rows.append([
            Cell(),
            Cell(text=f'ILO{index + 1}', align='center', bold=True),
            *[Cell(text=value, align='center') for value in values],
        ])
    data_width = 5.8 / len(headers)
    widths = [0.9, 0.5] + [data_width] * len(headers)
    return Table(rows=rows, widths=widths, header_rows=2, size=7.5, keep_together=True)


def _cdio_sdg_mapping_table(data):
    cdio_headers = _mapping_headers(data, 'cdio_headers', 4)
    sdg_headers = _mapping_headers(data, 'sdg_headers', 5)
    source_rows = _normalize_mapping(data, 'cdio_sdg_rows')
    rows = [[
        _label('ILO-CDIO and ILO-SDG Mapping', rowspan=2),
        _head('ILOs', rowspan=2),
        _head('CDIO SKILLS', colspan=len(cdio_headers), align='center'),
        _head('SDG Skills', colspan=len(sdg_headers), align='center'),
    ], [
        *[_head(header) for header in cdio_headers],
        *[_head(header) for header in sdg_headers],
    ]]
    for index in range(7):
        source = source_rows[index] if index < len(source_rows) else {}
        cdio_values = _pad(source.get('cdio_vals') or [], len(cdio_headers))
        sdg_values = _pad(source.get('sdg_vals') or [], len(sdg_headers))
        rows.append([
            Cell(),
            Cell(text=f'ILO{index + 1}', align='center', bold=True),
            *[Cell(text=value, align='center') for value in cdio_values],
            *[Cell(text=value, align='center') for value in sdg_values],
        ])
    data_columns = len(cdio_headers) + len(sdg_headers)
    data_width = 5.8 / data_columns if data_columns else 1.0
    widths = [0.9, 0.5] + [data_width] * data_columns
    return Table(rows=rows, widths=widths, header_rows=2, size=7.5, keep_together=True)


def _signature_table(data):
    prepared = _first(data, 'instructor_name')
    reviewed = _first(data, 'reviewed_by')
    approved = _first(data, 'approved_by')
    rows = [
        [Cell(text='Prepared by:', bold=True), Cell(text='Reviewed by:', bold=True), Cell(text='Approved by:', bold=True)],
        [Cell(), Cell(), Cell()],
        [
            Cell(text=prepared, bold=True, align='center'),
            Cell(text=reviewed, bold=True, align='center'),
            Cell(text=approved, bold=True, align='center'),
        ],
        [
            Cell(text=_first(data, 'prepared_by_title') or 'Faculty', align='center'),
            Cell(text=_first(data, 'reviewed_by_title'), align='center'),
            Cell(text=_first(data, 'approved_by_title'), align='center'),
        ],
        [Cell(text='Date:'), Cell(text='Date:'), Cell(text='Date:')],
    ]
    return Table(rows=rows, widths=[2.4, 2.4, 2.4], size=8.5, keep_together=True)


def build_standard_document(syllabus, branding=None):
    data = syllabus if isinstance(syllabus, dict) else {}
    branding = branding if isinstance(branding, dict) else {}
    blocks = [
        _course_info_table(data),
        _ilo_table(data),
        _at_table(data),
        _references_table(data),
        _framework_table(
            data, 'igas', 'IGA', 8,
            'Institutional Graduate Attributes (IGA)', 'IGA',
            'Institutional Graduate Attributes (IGA) Statements',
        ),
        _framework_table(
            data, 'student_outcomes', 'SO', 6,
            'Student Outcomes (SO)', 'SO', 'Student Outcomes (SO) Statements',
        ),
        _framework_table(
            data, 'cdio_skills', 'CDIO', 4,
            'CDIO Framework Skills', 'CDIO', 'CDIO Skills',
        ),
        _framework_table(
            data, 'sdg_skills', 'SDG', 5,
            'Sustainable Development Goals Skills', 'SDG', 'SDG Skills',
        ),
        _policies_table(data),
        _tla_table(data),
        _schedule_table(data),
        _so_mapping_table(data),
        _iga_mapping_table(data),
        _cdio_sdg_mapping_table(data),
        _signature_table(data),
        Paragraph(text='Remarks:', bold=True, size=8.5),
    ]
    for index, remark in enumerate(REMARKS, start=1):
        blocks.append(Paragraph(text=f'{index}  {remark}', size=8))
    return StandardDocument(
        letterhead=_letterhead(data, branding),
        blocks=blocks,
        college=FIXED_COLLEGE,
        program=FIXED_PROGRAM,
        title=FIXED_TITLE,
    )


# ── shared text helpers ─────────────────────────────────────────────────────


def _decode_image(value):
    if not isinstance(value, str):
        return None
    value = value.strip()
    if not value.startswith('data:image'):
        return None
    match = re.match(r'data:image/[^;]+;base64,(.*)', value, re.S)
    if not match:
        return None
    try:
        return base64.b64decode(match.group(1), validate=False)
    except (binascii.Error, ValueError):
        return None


def _line_markup(cell):
    """ReportLab-flavoured markup for a cell's lines."""
    parts = []
    for index, line in enumerate((cell.text or '').split('\n')):
        text = xml_escape(line)
        if cell.bold or (cell.first_line_bold and index == 0):
            text = f'<b>{text}</b>'
        if cell.italic or (cell.rest_italic and index > 0):
            text = f'<i>{text}</i>'
        parts.append(text)
    return '<br/>'.join(parts)


# ── DOCX renderer ───────────────────────────────────────────────────────────


def _add_bottom_border(paragraph):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '8')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '000000')
    borders.append(bottom)
    properties.insert_element_before(
        borders,
        'w:shd', 'w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap',
        'w:overflowPunct', 'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN',
        'w:bidi', 'w:adjustRightInd', 'w:snapToGrid', 'w:spacing', 'w:ind',
        'w:contextualSpacing', 'w:mirrorIndents', 'w:suppressOverlap', 'w:jc',
        'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap',
        'w:outlineLvl', 'w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr',
        'w:pPrChange',
    )


def _shade_cell(cell, fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = cell._tc.get_or_add_tcPr()
    for existing in properties.findall(qn('w:shd')):
        properties.remove(existing)
    shading = OxmlElement('w:shd')
    shading.set(qn('w:val'), 'clear')
    shading.set(qn('w:color'), 'auto')
    shading.set(qn('w:fill'), fill)
    properties.insert_element_before(
        shading,
        'w:noWrap', 'w:tcMar', 'w:textDirection', 'w:tcFitText', 'w:vAlign',
        'w:hideMark',
    )


def _docx_alignment(align):
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    return {
        'center': WD_ALIGN_PARAGRAPH.CENTER,
        'right': WD_ALIGN_PARAGRAPH.RIGHT,
    }.get(align, WD_ALIGN_PARAGRAPH.LEFT)


DOCX_FONT_NAME = 'Times New Roman'
PDF_FONT = 'Times-Roman'
PDF_FONT_BOLD = 'Times-Bold'


def _set_run_font_docx(run, size=None, bold=None, italic=None, color=None):
    """Pin every run to Times New Roman so Word matches the PDF/HTML metrics."""
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = DOCX_FONT_NAME
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = rpr.get_or_add_rFonts()
    for attribute in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(attribute), DOCX_FONT_NAME)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _add_letterhead_lines_docx(container, lines):
    from docx.shared import Pt

    for text, style in lines:
        paragraph = container.add_paragraph()
        paragraph.alignment = _docx_alignment('center')
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text)
        _set_run_font_docx(
            run,
            size=style.get('size', 8.5),
            bold=bool(style.get('bold')),
            color=style.get('color'),
        )


CONTENT_WIDTH_INCHES = 7.27
# A portrait-shaped upload must not push the rest of the syllabus off the page.
MAX_HEADER_INCHES = 2.5


def _header_image_size(px_width, px_height):
    """Fit the uploaded template to the text column, capped in height."""
    if not px_width or not px_height:
        return CONTENT_WIDTH_INCHES, MAX_HEADER_INCHES
    width = CONTENT_WIDTH_INCHES
    height = width * px_height / px_width
    if height > MAX_HEADER_INCHES:
        height = MAX_HEADER_INCHES
        width = height * px_width / px_height
    return width, height


def _render_letterhead_docx(doc, letterhead):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.image.image import Image as DocxImage
    from docx.shared import Inches, Pt

    image = _decode_image(letterhead.image)
    if image:
        # One header template; it is the whole letterhead, scaled to the text column.
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        try:
            info = DocxImage.from_blob(image)
            width, height = _header_image_size(info.px_width, info.px_height)
            paragraph.add_run().add_picture(
                BytesIO(image),
                width=Inches(width),
                height=Inches(height),
            )
        except Exception:
            _add_letterhead_lines_docx(doc, letterhead.lines)
    else:
        _add_letterhead_lines_docx(doc, letterhead.lines)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(0)
    rule.paragraph_format.space_after = Pt(2)
    _add_bottom_border(rule)


def _add_heading_lines_docx(doc, document):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    if document.college:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(document.college)
        _set_run_font_docx(run, size=10, bold=True)

    if document.program:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(document.program)
        _set_run_font_docx(run, size=9.5, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(2)
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run(document.title)
    _set_run_font_docx(run, size=11, bold=True)


def _add_paragraph_docx(doc, block):
    from docx.shared import Pt

    paragraph = doc.add_paragraph()
    paragraph.alignment = _docx_alignment(block.align)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(1.5)
    run = paragraph.add_run(block.text)
    _set_run_font_docx(run, size=block.size, bold=block.bold, italic=block.italic)


def _fill_cell_docx(cell, data_cell, default_size):
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.shared import Pt

    cell.text = ''
    paragraph = cell.paragraphs[0]
    for index, line in enumerate((data_cell.text or '').split('\n')):
        if index:
            paragraph = cell.add_paragraph()
        paragraph.alignment = _docx_alignment(data_cell.align)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(line)
        _set_run_font_docx(
            run,
            size=default_size,
            bold=data_cell.bold or (data_cell.first_line_bold and index == 0),
            italic=data_cell.italic or (data_cell.rest_italic and index > 0),
        )
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    if data_cell.shade:
        _shade_cell(cell, HEADER_FILL)


def _repeat_row_docx(row):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement('w:tblHeader')
    marker.set(qn('w:val'), 'true')
    properties.insert_element_before(
        marker, 'w:tblCellSpacing', 'w:jc', 'w:hidden', 'w:ins', 'w:del',
        'w:trPrChange',
    )


def _no_split_row_docx(row):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = row._tr.get_or_add_trPr()
    marker = OxmlElement('w:cantSplit')
    marker.set(qn('w:val'), 'true')
    properties.insert_element_before(
        marker, 'w:trHeight', 'w:tblHeader', 'w:tblCellSpacing', 'w:jc', 'w:hidden',
        'w:ins', 'w:del', 'w:trPrChange',
    )


def _set_style_font(style, name=DOCX_FONT_NAME, size=8.5):
    from docx.oxml.ns import qn
    from docx.shared import Pt

    style.font.name = name
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attribute in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(attribute), name)


def _compact_table_docx(table, widths, *, content_width=7.27, margins=(2, 2, 1, 1)):
    """Force fixed layout, full content width, and tight cell margins.

    Word's default table metrics are much roomier than the CIS form, so the
    widths of the shared document model would otherwise be ignored and rows
    would grow taller than the reference.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches

    properties = table._tbl.tblPr

    for tag in ('w:tblW', 'w:tblCellMar'):
        for existing in properties.findall(qn(tag)):
            properties.remove(existing)

    width = OxmlElement('w:tblW')
    width.set(qn('w:w'), '5000')
    width.set(qn('w:type'), 'pct')
    properties.insert_element_before(
        width,
        'w:jc', 'w:tblCellSpacing', 'w:tblInd', 'w:tblBorders', 'w:shd',
        'w:tblLayout', 'w:tblCellMar', 'w:tblLook',
    )

    left, right, top, bottom = margins
    margin = OxmlElement('w:tblCellMar')
    for tag, value in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        element = OxmlElement(f'w:{tag}')
        element.set(qn('w:w'), str(int(value * 20)))
        element.set(qn('w:type'), 'dxa')
        margin.append(element)
    properties.insert_element_before(margin, 'w:tblLook')

    total = sum(widths) or 1.0
    scale = content_width / total
    for index, value in enumerate(widths):
        if index < len(table.columns):
            table.columns[index].width = Inches(value * scale)


def _hide_label_borders_docx(cell, *, first, last):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for tag in ('w:top', 'w:left', 'w:bottom', 'w:right'):
        hide = (tag == 'w:top' and not first) or (tag == 'w:bottom' and not last)
        if not hide:
            continue
        element = OxmlElement(tag)
        element.set(qn('w:val'), 'nil')
        borders.append(element)
    tc_pr.insert_element_before(
        borders, 'w:shd', 'w:noWrap', 'w:tcMar', 'w:textDirection',
        'w:tcFitText', 'w:vAlign', 'w:hideMark',
    )


def _apply_cell_widths_docx(table, inch_widths):
    """Pin every cell to its declared width, merged cells included.

    python-docx only writes the grid column widths; Word then re-fits merged
    cells from their content, which drifts away from the CIS form. Writing an
    explicit ``w:tcW`` per cell (summed across its span) keeps both renderers
    on the same column geometry.
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    for row in table.rows:
        position = 0
        for tc in row._tr.tc_lst:
            tc_pr = tc.get_or_add_tcPr()
            span = 1
            grid_span = tc_pr.find(qn('w:gridSpan'))
            if grid_span is not None:
                span = int(grid_span.get(qn('w:val')))
            width = sum(inch_widths[position:position + span])
            for existing in tc_pr.findall(qn('w:tcW')):
                tc_pr.remove(existing)
            element = OxmlElement('w:tcW')
            element.set(qn('w:w'), str(int(round(width * 1440))))
            element.set(qn('w:type'), 'dxa')
            tc_pr.insert(0, element)
            position += span


def _add_table_docx(doc, block):
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Inches, Pt

    rows = block.rows
    columns = _table_columns(rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [
        block.widths[index] if index < len(block.widths) else 1.0
        for index in range(columns)
    ]
    _compact_table_docx(table, widths)

    occupied = set()
    for row_index, row in enumerate(rows):
        column = 0
        for data_cell in row:
            while (row_index, column) in occupied:
                column += 1
            stacked_span = (
                block.keep_together
                and data_cell.colspan == 1
                and data_cell.rowspan > 1
            )
            if stacked_span:
                # Any vertical merge makes Word ignore keep-with-next, so a
                # spanning label forces the form block to break across pages and
                # clips the text. Emit the same look with stacked cells whose
                # shared horizontal edges are hidden.
                for offset in range(data_cell.rowspan):
                    occupied.add((row_index + offset, column))
                    stacked = table.cell(row_index + offset, column)
                    _fill_cell_docx(
                        stacked,
                        data_cell if offset == 0 else Cell(),
                        block.size,
                    )
                    _hide_label_borders_docx(
                        stacked,
                        first=offset == 0,
                        last=offset == data_cell.rowspan - 1,
                    )
                column += 1
                continue
            if data_cell.rowspan > 1 or data_cell.colspan > 1:
                for span_row in range(row_index, row_index + data_cell.rowspan):
                    for span_col in range(column, column + data_cell.colspan):
                        occupied.add((span_row, span_col))
                target = table.cell(row_index, column).merge(
                    table.cell(
                        row_index + data_cell.rowspan - 1,
                        column + data_cell.colspan - 1,
                    )
                )
            else:
                occupied.add((row_index, column))
                target = table.cell(row_index, column)
            _fill_cell_docx(target, data_cell, block.size)
            column += data_cell.colspan

    total = sum(widths) or 1.0
    _apply_cell_widths_docx(table, [value * 7.27 / total for value in widths])

    for header_index in range(min(block.header_rows, len(table.rows))):
        _repeat_row_docx(table.rows[header_index])

    if block.no_split or block.keep_together:
        for row in table.rows:
            _no_split_row_docx(row)

    if block.keep_together:
        # keep_with_next chains the rows: Word moves the whole form block to the
        # next page instead of clipping a merged label cell across a page break.
        for row in table.rows[:-1]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.add_run('').font.size = Pt(1)


def _add_field(paragraph, instruction):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def _field_run(element):
        run = OxmlElement('w:r')
        run.append(element)
        paragraph._p.append(run)

    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    _field_run(begin)
    text = OxmlElement('w:instrText')
    text.set(qn('xml:space'), 'preserve')
    text.text = instruction
    _field_run(text)
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    _field_run(separate)
    result = OxmlElement('w:t')
    result.text = '1'
    _field_run(result)
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    _field_run(end)


def _add_page_footer(doc):
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    for section in doc.sections:
        footer = section.footer
        footer.is_linked_to_previous = False
        paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run('Page ')
        _set_run_font_docx(run, size=7.5)
        _add_field(paragraph, ' PAGE ')
        run = paragraph.add_run(' of ')
        _set_run_font_docx(run, size=7.5)
        _add_field(paragraph, ' NUMPAGES ')


def build_standard_syllabus_docx(syllabus, branding=None):
    from docx import Document
    from docx.shared import Inches

    document = build_standard_document(syllabus, branding)
    doc = Document()
    _set_style_font(doc.styles['Normal'])
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    _render_letterhead_docx(doc, document.letterhead)
    _add_heading_lines_docx(doc, document)
    for block in document.blocks:
        if isinstance(block, Paragraph):
            _add_paragraph_docx(doc, block)
        else:
            _add_table_docx(doc, block)
    _add_page_footer(doc)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ── HTML / WeasyPrint renderer ──────────────────────────────────────────────


_HTML_CSS = """
@page {
  size: A4;
  margin: 12mm 12mm 14mm 12mm;
  @bottom-right {
    content: "Page " counter(page) " of " counter(pages);
    font-family: "Times New Roman", Times, serif;
    font-size: 7.5pt;
    color: #444444;
  }
}
:root { --accent: {accent}; }
* { box-sizing: border-box; }
body {{
  font-family: "Times New Roman", Times, serif;
  font-size: 8.5pt;
  line-height: 1.25;
  color: #000000;
  margin: 0;
}}
.letterhead {{ width: 100%; border-collapse: collapse; table-layout: fixed; }}
.letterhead td {{ border: none; padding: 0; vertical-align: middle; }}
.letterhead-banner {{ display: block; width: 100%; height: auto; max-height: 2.5in; object-fit: contain; margin: 0 auto; }}
.letterhead-text {{ text-align: center; }}
.letterhead-text div {{ margin: 0; }}
.rule {{ border-bottom: 0.8pt solid #000; margin: 0 0 3pt; }}
.college {{ font-size: 10pt; font-weight: bold; margin: 2pt 0; }}
.program {{ font-size: 9.5pt; font-weight: bold; text-align: center; margin: 2pt 0; }}
.doc-title {{ font-size: 11pt; font-weight: bold; text-align: center; margin: 3pt 0 5pt; }}
table {{ width: 100%; border-collapse: collapse; table-layout: fixed; margin: 0 0 1.5pt; }}
td {{ border: 0.75pt solid #000; padding: 1.5pt 3pt; vertical-align: top; word-wrap: break-word; }}
td.head {{ background: #f2f2f2; font-weight: bold; text-align: center; }}
p {{ margin: 0 0 1.5pt; }}
"""


def _html_cell(cell):
    parts = []
    for index, line in enumerate((cell.text or '').split('\n')):
        text = html.escape(line)
        if cell.bold or (cell.first_line_bold and index == 0):
            text = f'<b>{text}</b>'
        if cell.italic or (cell.rest_italic and index > 0):
            text = f'<i>{text}</i>'
        parts.append(text)
    return '<br/>'.join(parts)


def _render_letterhead_html(letterhead):
    if letterhead.image:
        # The admin's uploaded template is the whole letterhead; no text sits beside it.
        return (
            f'<img class="letterhead-banner" src="{html.escape(letterhead.image, quote=True)}" '
            'alt="Letterhead"/>'
            '<div class="rule"></div>'
        )
    lines = []
    for text, style in letterhead.lines:
        css = [f"font-size:{style.get('size', 8.5)}pt"]
        if style.get('bold'):
            css.append('font-weight:bold')
        if style.get('color'):
            css.append(f"color:#{style['color']}")
        lines.append(f'<div style="{";".join(css)}">{html.escape(text)}</div>')
    return (
        '<table class="letterhead"><tr>'
        f'<td class="letterhead-text">{"".join(lines)}</td>'
        '</tr></table>'
        '<div class="rule"></div>'
    )


def _render_table_html(block):
    rows = block.rows
    columns = _table_columns(rows)
    widths = list(block.widths[:columns])
    while len(widths) < columns:
        widths.append(1.0)
    total = sum(widths) or 1.0
    colgroup = ''.join(
        f'<col style="width:{width / total * 100:.3f}%"/>' for width in widths
    )
    occupied = set()
    body = []
    for row_index, row in enumerate(rows):
        column = 0
        cells = []
        for data_cell in row:
            while (row_index, column) in occupied:
                column += 1
            attributes = []
            if data_cell.colspan > 1:
                attributes.append(f'colspan="{data_cell.colspan}"')
            if data_cell.rowspan > 1:
                attributes.append(f'rowspan="{data_cell.rowspan}"')
            classes = ' class="head"' if data_cell.shade else ''
            style = f"text-align:{data_cell.align};"
            if data_cell.bold:
                style += 'font-weight:bold;'
            if data_cell.italic:
                style += 'font-style:italic;'
            cells.append(
                f'<td{" " + " ".join(attributes) if attributes else ""}{classes} '
                f'style="{style}">{_html_cell(data_cell)}</td>'
            )
            for span_row in range(row_index, row_index + data_cell.rowspan):
                for span_col in range(column, column + data_cell.colspan):
                    occupied.add((span_row, span_col))
            column += data_cell.colspan
        body.append(f'<tr>{"".join(cells)}</tr>')
    return f'<table><colgroup>{colgroup}</colgroup><tbody>{"".join(body)}</tbody></table>'


def build_standard_syllabus_html(syllabus, branding=None, style_profile=None):
    branding = branding if isinstance(branding, dict) else {}
    style_profile = style_profile if isinstance(style_profile, dict) else (
        branding.get('syllabus_style_profile') or {}
    )
    accent = ((style_profile.get('colors') or {}).get('accent') or '#003366')
    if not str(accent).startswith('#'):
        accent = '#' + str(accent)
    document = build_standard_document(syllabus, branding)
    parts = [
        '<!DOCTYPE html><html><head><meta charset="utf-8"/>',
        f'<title>{html.escape(document.title)}</title>',
        '<style>',
        _HTML_CSS.replace('{accent}', accent),
        '</style></head><body>',
        _render_letterhead_html(document.letterhead),
    ]
    if document.college:
        parts.append(f'<div class="college">{html.escape(document.college)}</div>')
    if document.program:
        parts.append(f'<div class="program">{html.escape(document.program)}</div>')
    parts.append(f'<div class="doc-title">{html.escape(document.title)}</div>')
    for block in document.blocks:
        if isinstance(block, Paragraph):
            css = [f"font-size:{block.size}pt"]
            if block.bold:
                css.append('font-weight:bold')
            if block.italic:
                css.append('font-style:italic')
            parts.append(
                f'<p style="{";".join(css)};text-align:{block.align}">{html.escape(block.text)}</p>'
            )
        else:
            parts.append(_render_table_html(block))
    parts.append('</body></html>')
    return ''.join(parts)


# ── ReportLab renderer ──────────────────────────────────────────────────────


def _build_reportlab_story(document):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import Image, Paragraph as RLParagraph, Spacer, Table as RLTable, TableStyle
    from reportlab.platypus.flowables import HRFlowable

    alignment = {'left': TA_LEFT, 'center': TA_CENTER, 'right': TA_RIGHT}
    styles = {}

    def paragraph_style(size, align):
        key = (size, align)
        if key not in styles:
            styles[key] = ParagraphStyle(
                f'cis-{size}-{align}',
                fontName=PDF_FONT,
                fontSize=size,
                leading=size * 1.18,
                alignment=alignment.get(align, TA_LEFT),
                spaceBefore=0,
                spaceAfter=0,
            )
        return styles[key]

    story = []

    letterhead = document.letterhead
    image_bytes = _decode_image(letterhead.image)
    header_image = None
    if image_bytes:
        try:
            header_image = Image(BytesIO(image_bytes))
            if header_image.imageWidth and header_image.imageHeight:
                width, height = _header_image_size(
                    header_image.imageWidth, header_image.imageHeight,
                )
                header_image.drawWidth = width * inch
                header_image.drawHeight = height * inch
                header_image.hAlign = 'CENTER'
            else:
                header_image = None
        except Exception:
            header_image = None

    line_flowables = []
    for text, style in letterhead.lines:
        paragraph = RLParagraph(
            xml_escape(text),
            ParagraphStyle(
                f'lh-{style.get("size", 8.5)}-{bool(style.get("bold"))}',
                fontName=PDF_FONT_BOLD if style.get('bold') else PDF_FONT,
                fontSize=style.get('size', 8.5),
                leading=style.get('size', 8.5) * 1.15,
                alignment=TA_CENTER,
                textColor=colors.HexColor(f'#{style["color"]}') if style.get('color') else colors.black,
            ),
        )
        line_flowables.append(paragraph)

    if header_image is not None:
        # The uploaded template is the entire letterhead — nothing is layered beside it.
        story.append(header_image)
    elif line_flowables:
        header_table = RLTable(
            [[line_flowables]],
            colWidths=[CONTENT_WIDTH_INCHES * inch],
        )
        header_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('ALIGN', (0, 0), (0, 0), 'CENTER'),
            ('LEFTPADDING', (0, 0), (-1, -1), 0),
            ('RIGHTPADDING', (0, 0), (-1, -1), 0),
            ('TOPPADDING', (0, 0), (-1, -1), 0),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
        ]))
        story.append(header_table)
    story.append(HRFlowable(width='100%', thickness=0.8, color=colors.black, spaceBefore=1, spaceAfter=2))

    if document.college:
        story.append(RLParagraph(
            xml_escape(document.college),
            ParagraphStyle('college', fontName=PDF_FONT_BOLD, fontSize=10, leading=12, spaceBefore=2, spaceAfter=2),
        ))
    if document.program:
        story.append(RLParagraph(
            xml_escape(document.program),
            ParagraphStyle('program', fontName=PDF_FONT_BOLD, fontSize=9.5, leading=11.5, alignment=TA_CENTER, spaceBefore=2, spaceAfter=2),
        ))
    story.append(RLParagraph(
        xml_escape(document.title),
        ParagraphStyle('doc-title', fontName=PDF_FONT_BOLD, fontSize=11, leading=13, alignment=TA_CENTER, spaceBefore=2, spaceAfter=4),
    ))

    def table_flowable(block):
        rows = block.rows
        columns = _table_columns(rows)
        widths = list(block.widths[:columns])
        while len(widths) < columns:
            widths.append(1.0)
        total = sum(widths) or 1.0
        available = 7.27 * inch
        scale = available / (total * inch) if total * inch > available else 1.0
        col_widths = [width * inch * scale for width in widths]

        grid = [['' for _ in range(columns)] for _ in rows]
        spans = []
        backgrounds = []
        occupied = set()
        for row_index, row in enumerate(rows):
            column = 0
            for data_cell in row:
                while (row_index, column) in occupied:
                    column += 1
                markup = _line_markup(data_cell)
                grid[row_index][column] = RLParagraph(
                    markup,
                    paragraph_style(block.size, data_cell.align),
                ) if markup else ''
                if data_cell.rowspan > 1 or data_cell.colspan > 1:
                    spans.append((
                        'SPAN',
                        (column, row_index),
                        (column + data_cell.colspan - 1, row_index + data_cell.rowspan - 1),
                    ))
                if data_cell.shade:
                    backgrounds.append((
                        'BACKGROUND',
                        (column, row_index),
                        (column + data_cell.colspan - 1, row_index + data_cell.rowspan - 1),
                        colors.HexColor('#F2F2F2'),
                    ))
                for span_row in range(row_index, row_index + data_cell.rowspan):
                    for span_col in range(column, column + data_cell.colspan):
                        occupied.add((span_row, span_col))
                column += data_cell.colspan

        table = RLTable(grid, colWidths=col_widths, repeatRows=block.header_rows)
        commands = [
            ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 2.5),
            ('RIGHTPADDING', (0, 0), (-1, -1), 2.5),
            ('TOPPADDING', (0, 0), (-1, -1), 1.2),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 1.2),
            *backgrounds,
            *spans,
        ]
        table.setStyle(TableStyle(commands))
        return table

    for block in document.blocks:
        if isinstance(block, Paragraph):
            story.append(RLParagraph(
                xml_escape(block.text),
                ParagraphStyle(
                    f'para-{block.size}-{block.bold}',
                    fontName=PDF_FONT_BOLD if block.bold else PDF_FONT,
                    fontSize=block.size,
                    leading=block.size * 1.2,
                    alignment=alignment.get(block.align, TA_LEFT),
                    spaceBefore=0,
                    spaceAfter=1.5,
                ),
            ))
        else:
            story.append(table_flowable(block))
            story.append(Spacer(1, 1.5))
    return story


def _reportlab_canvas():
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas as canvas_module

    class NumberedCanvas(canvas_module.Canvas):
        def __init__(self, *args, **kwargs):
            canvas_module.Canvas.__init__(self, *args, **kwargs)
            self._saved_page_states = []

        def showPage(self):
            self._saved_page_states.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._saved_page_states)
            for state in self._saved_page_states:
                self.__dict__.update(state)
                self.setFont(PDF_FONT, 7.5)
                self.setFillColorRGB(0.27, 0.27, 0.27)
                self.drawRightString(A4[0] - 34, 18, f'Page {self._pageNumber} of {total}')
                canvas_module.Canvas.showPage(self)
            canvas_module.Canvas.save(self)

    return NumberedCanvas


def _pdf_via_reportlab(document):
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        topMargin=0.5 * 72,
        bottomMargin=0.5 * 72,
        leftMargin=0.5 * 72,
        rightMargin=0.5 * 72,
        title=document.title,
    )
    doc.build(
        _build_reportlab_story(document),
        canvasmaker=_reportlab_canvas(),
    )
    return buffer.getvalue()


def build_standard_syllabus_pdf(syllabus, branding=None, style_profile=None):
    """Return (pdf_bytes, engine) for the fixed CIS layout.

    ReportLab is used directly: it renders the same document model as the DOCX
    export, so the PDF and DOCX cannot drift apart. WeasyPrint is deliberately
    not used as the primary path because its HTML layout differs in font
    metrics, cell padding and page breaks from the Word output.
    """
    document = build_standard_document(syllabus, branding)
    return _pdf_via_reportlab(document), 'reportlab'
