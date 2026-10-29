from types import SimpleNamespace
from unittest.mock import patch

from django.test import SimpleTestCase
from rest_framework.test import APIRequestFactory, force_authenticate

from apps.curriculum.docx_builders import build_tos_docx
from apps.curriculum.views import (
    _align_generated_section_to_blueprint,
    _build_plain_assessment_docx,
    _build_sectioned_assessment_payload,
    _coerce_generated_item_fields,
    _collect_section_items,
    _compact_assessment_slot,
    _derive_assessment_source,
    _fill_answer_key_fallback,
    _generated_item_quality_errors,
    _materialize_tos_exam,
    _normalize_slot_item,
    create_course,
)
from apps.integrations.syllabus_json_import import syllabus_content_from_markdown
from apps.integrations.syllabus_tab_split import (
    faithful_parsed_data,
    split_markdown_into_tabs,
    supplement_markdown,
)


class FacultyCourseInstitutionScopeTests(SimpleTestCase):
    def setUp(self):
        self.factory = APIRequestFactory()

    @patch('apps.curriculum.views._is_faculty_user', return_value=True)
    @patch('apps.curriculum.views.Course.objects.filter')
    @patch('apps.curriculum.views.Course.objects.create')
    @patch('apps.curriculum.views.CourseLoad.objects.create')
    def test_created_course_and_assignment_belong_to_faculty_institution(
        self,
        create_assignment,
        create_course_record,
        filter_courses,
        _is_faculty,
    ):
        user = SimpleNamespace(
            is_authenticated=True,
            institution_id='institution-a',
            email='faculty@example.com',
        )
        course = SimpleNamespace(
            course_id='course-a',
            course_code='CS101',
            course_title='Computer Science',
            description=None,
        )
        filter_courses.return_value.exists.return_value = False
        create_course_record.return_value = course
        request = self.factory.post(
            '/api/curriculum/courses/',
            {'course_code': 'CS101', 'course_title': 'Computer Science'},
            format='json',
        )
        force_authenticate(request, user=user)

        response = create_course(request)

        self.assertEqual(response.status_code, 201)
        filter_courses.assert_called_once_with(
            course_code='CS101',
            institution_id='institution-a',
        )
        create_course_record.assert_called_once_with(
            institution_id='institution-a',
            course_code='CS101',
            course_title='Computer Science',
            description=None,
        )
        create_assignment.assert_called_once_with(course=course, faculty=user)
        self.assertTrue(_is_faculty.called)

    @patch('apps.curriculum.views._is_faculty_user', return_value=True)
    @patch('apps.curriculum.views.Course.objects.filter')
    @patch('apps.curriculum.views.Course.objects.create')
    def test_course_creation_requires_an_institution(
        self,
        create_course_record,
        filter_courses,
        _is_faculty,
    ):
        user = SimpleNamespace(
            is_authenticated=True,
            institution_id=None,
            email='faculty@example.com',
        )
        request = self.factory.post(
            '/api/curriculum/courses/',
            {'course_code': 'CS101', 'course_title': 'Computer Science'},
            format='json',
        )
        force_authenticate(request, user=user)

        response = create_course(request)

        self.assertEqual(response.status_code, 400)
        filter_courses.assert_not_called()
        create_course_record.assert_not_called()
        self.assertTrue(_is_faculty.called)


class TosFormulaTests(SimpleTestCase):
    def test_hours_split_points_minutes_and_bloom_counts_exactly(self):
        exam = _materialize_tos_exam({
            'totalTestItems': 10,
            'totalExamMinutes': 100,
            'topics': [
                {
                    'topic': 'Hardware',
                    'hours_class': 1,
                    'bloom_pct': {'remember': 50, 'understand': 50},
                },
                {
                    'topic': 'Operating systems',
                    'hours_class': 1,
                    'bloom_pct': {'apply': 100},
                },
                {
                    'topic': 'Access control',
                    'hours_class': 2,
                    'bloom_pct': {'analyze': 100},
                },
            ],
        })

        topics = exam['topics']
        self.assertEqual([topic['total_points'] for topic in topics], [3, 2, 5])
        self.assertEqual([topic['hours_min'] for topic in topics], [25, 25, 50])
        self.assertEqual(topics[0]['bloom_items']['remember'], 2)
        self.assertEqual(topics[0]['bloom_items']['understand'], 1)
        self.assertEqual(topics[0]['bloom_items']['apply'], 0)
        self.assertEqual(topics[1]['bloom_items']['apply'], 2)
        self.assertEqual(topics[2]['bloom_items']['analyze'], 5)
        self.assertEqual(sum(topic['total_points'] for topic in topics), 10)
        self.assertEqual(sum(topic['hours_min'] for topic in topics), 100)

    def test_assessment_items_keep_tos_bloom_counts(self):
        exam = _materialize_tos_exam({
            'totalTestItems': 10,
            'topics': [
                {
                    'topic': 'Hardware',
                    'hours_class': 1,
                    'bloom_pct': {'remember': 50, 'understand': 50},
                },
                {
                    'topic': 'Operating systems',
                    'hours_class': 1,
                    'bloom_pct': {'apply': 100},
                },
            ],
        })
        sections = _build_sectioned_assessment_payload(
            exam['topics'],
            'major examination',
            99,
            assessment_label='Prelim',
        )
        counts = {}
        for section in sections:
            for item in section['items']:
                key = (section['section_title'], item['bloom_level'])
                counts[key] = counts.get(key, 0) + 1

        self.assertEqual(sum(len(section['items']) for section in sections), 10)
        self.assertEqual(counts[('Hardware', 'Remember')], 3)
        self.assertEqual(counts[('Hardware', 'Understand')], 2)
        self.assertEqual(counts[('Operating systems', 'Apply')], 5)
        self.assertNotIn(('Hardware', 'Apply'), counts)
        self.assertNotIn(('Hardware', 'Create'), counts)
        by_level = {}
        for section in sections:
            for item in section['items']:
                by_level.setdefault(item['bloom_level'], []).append(item['item_number'])
        self.assertEqual(by_level['Remember'], [1, 2, 3])
        self.assertEqual(by_level['Understand'], [4, 5])
        self.assertEqual(by_level['Apply'], [6, 7, 8, 9, 10])

    def test_checked_ilos_set_the_bloom_levels_and_numbering_stays_with_the_topic(self):
        sections = _build_sectioned_assessment_payload(
            [
                {'topic': 'Cells', 'total_points': 10, 'ilos': ['ILO1', 'ILO2']},
                {
                    'topic': 'Genetics',
                    'total_points': 4,
                    'bloom_items': {'remember': 1, 'create': 3},
                    'ilos': ['ILO2'],
                },
            ],
            'major examination',
            14,
            ilos=['Identify cell parts', 'Design an experiment'],
            ilo_bloom=[
                {'level': 'Remember', 'verb': 'identify'},
                {'level': 'Create', 'verb': 'design'},
            ],
            assessment_label='Prelim',
        )
        cells = sections[0]['items']
        genetics = sections[1]['items']
        self.assertEqual(
            [item['bloom_level'] for item in cells],
            ['Remember'] * 5 + ['Create'] * 5,
        )
        self.assertEqual(cells[0]['bloom_verb'], 'identify')
        self.assertEqual(cells[5]['bloom_verb'], 'design')
        self.assertEqual([item['item_number'] for item in cells], list(range(1, 11)))
        self.assertEqual(
            [item['bloom_level'] for item in genetics],
            ['Remember'] + ['Create'] * 3,
        )
        self.assertEqual([item['item_number'] for item in genetics], [11, 12, 13, 14])
        self.assertEqual(genetics[1]['bloom_verb'], 'design')

    def test_hours_only_rows_stay_remember_until_faculty_set_a_level(self):
        sections = _build_sectioned_assessment_payload(
            [
                {'topic': 'Cells', 'total_points': 4},
                {'topic': 'Genetics', 'total_points': 4},
            ],
            'major examination',
            8,
            assessment_label='Final',
        )
        levels = [item['bloom_level'] for section in sections for item in section['items']]
        self.assertEqual(levels, ['Remember'] * 8)
        self.assertEqual(
            [item['item_number'] for section in sections for item in section['items']],
            list(range(1, 9)),
        )

    def test_generated_items_are_forced_back_to_blueprint_blooms(self):
        blueprint = {
            'section_title': 'Hardware',
            'exam_label': 'Prelim',
            'bloom_distribution': {'Remember': 2, 'Apply': 1},
            'items': [
                {'item_number': 1, 'bloom_level': 'Remember', 'bloom_verb': 'list', 'ilo_refs': ['ILO1'], 'question_text': 'slot 1'},
                {'item_number': 2, 'bloom_level': 'Remember', 'bloom_verb': 'list', 'ilo_refs': ['ILO1'], 'question_text': 'slot 2'},
                {'item_number': 3, 'bloom_level': 'Apply', 'bloom_verb': 'use', 'ilo_refs': ['ILO2'], 'question_text': 'slot 3'},
            ],
        }
        generated = {
            'section_title': 'Hardware',
            'items': [
                {'item_number': 1, 'bloom_level': 'Create', 'question_text': 'Question one', 'choices': ['A. a', 'B. b', 'C. c', 'D. d'], 'correct_answer': 'B'},
                {'item_number': 2, 'bloom_level': 'Evaluate', 'question_text': 'Question two', 'choices': ['A. a', 'B. b', 'C. c', 'D. d'], 'correct_answer': 'C'},
                {'item_number': 3, 'bloom_level': 'Analyze', 'question_text': 'Question three', 'choices': ['A. a', 'B. b', 'C. c', 'D. d'], 'correct_answer': 'A'},
            ],
        }

        aligned = _align_generated_section_to_blueprint(generated, blueprint)
        self.assertEqual([item['bloom_level'] for item in aligned['items']], ['Remember', 'Remember', 'Apply'])
        self.assertEqual([item['question_text'] for item in aligned['items']], ['Question one', 'Question two', 'Question three'])
        self.assertEqual(aligned['allocated_items'], 3)


class AssessmentGenerationSlotTests(SimpleTestCase):
    def test_tos_exam_is_selected_by_exam_code_and_label(self):
        source = _derive_assessment_source(
            {
                'tos': [
                    {
                        'examCode': 'ME',
                        'examLabel': 'Midterm Examination',
                        'totalTestItems': 2,
                        'topics': [{'topic': 'Hardware', 'hours_class': 1, 'bloom_pct': {'remember': 100}}],
                    },
                    {
                        'examCode': 'FE',
                        'examLabel': 'Final Examination',
                        'totalTestItems': 2,
                        'topics': [{'topic': 'Networks', 'hours_class': 1, 'bloom_pct': {'apply': 100}}],
                    },
                ],
            },
            source_exam_code='FE',
            source_exam_label='Final Examination',
        )

        self.assertEqual(source['assessment_name'], 'Final Examination')
        self.assertEqual(source['topic_rows'][0]['topic'], 'Networks')

    def test_slot_prompt_omits_placeholder_stems_and_choices(self):
        slot = _compact_assessment_slot({
            'section_title': 'Hardware',
            'exam_label': 'Final Examination',
            'items': [{
                'item_number': 1,
                'question_text': '[Hardware] Major Examination item 1 (Remember)',
                'choices': ['A. ...', 'B. ...', 'C. ...', 'D. ...'],
                'bloom_level': 'Remember',
                'bloom_verb': 'list',
                'ilo_refs': ['ILO 1'],
            }],
        })

        self.assertNotIn('question_text', slot['items'][0])
        self.assertNotIn('choices', slot['items'][0])
        self.assertEqual(slot['items'][0]['bloom_level'], 'Remember')

    def test_explanation_alias_and_valid_item_are_kept(self):
        blueprint = {
            'section_title': 'Networks',
            'exam_label': 'Final Examination',
            'items': [
                {'item_number': 1, 'bloom_level': 'Apply', 'bloom_verb': 'use', 'ilo_refs': ['ILO 1'], 'question_text': '[Networks] Major Examination item 1 (Apply)'},
                {'item_number': 2, 'bloom_level': 'Apply', 'bloom_verb': 'use', 'ilo_refs': ['ILO 1'], 'question_text': '[Networks] Major Examination item 2 (Apply)'},
            ],
        }
        generated = {
            'items': [
                {
                    'item_number': 1,
                    'question_text': 'Which device forwards packets between networks?',
                    'choices': ['A. router', 'B. monitor', 'C. keyboard', 'D. printer'],
                    'correct_answer': 'A',
                    'explanation': 'A router forwards packets.',
                },
                {
                    'item_number': 2,
                    'question_text': 'Which command shows the route table?',
                    'choices': ['A. ...', 'B. ...', 'C. ...', 'D. ...'],
                    'correct_answer': 'B',
                    'answer_key': 'placeholder choices',
                },
            ],
        }

        accepted, retry_numbers = _collect_section_items(blueprint, generated)

        self.assertEqual(list(accepted), [1])
        self.assertEqual(accepted[1]['answer_key'], 'A router forwards packets.')
        self.assertEqual(accepted[1]['bloom_level'], 'Apply')
        self.assertEqual(retry_numbers, [2])

    def test_missing_explanation_gets_a_fallback_after_the_item_is_otherwise_complete(self):
        item = _normalize_slot_item(
            {
                'question_text': 'Which layer adds an IP header?',
                'choices': ['A. network', 'B. physical', 'C. session', 'D. application'],
                'correct_answer': 'A',
            },
            {'item_number': 3, 'bloom_level': 'Remember', 'bloom_verb': 'identify', 'ilo_refs': []},
            {'section_title': 'Networks', 'exam_label': 'Final Examination'},
        )
        self.assertEqual(_generated_item_quality_errors(item), ['answer_key (explanation) is missing.'])
        filled = _fill_answer_key_fallback(item)
        self.assertEqual(_generated_item_quality_errors(filled), [])
        self.assertIn('Option A', filled['answer_key'])

    def test_blueprint_stem_is_not_accepted_as_a_question(self):
        errors = _generated_item_quality_errors(_coerce_generated_item_fields({
            'question_text': '[Hardware] Major Examination item 1 (Remember)',
            'choices': ['A. ram', 'B. rom', 'C. cpu', 'D. bus'],
            'correct_answer': 'A',
            'answer_key': 'RAM is memory.',
        }))
        self.assertIn('question_text is missing.', errors)


class SyllabusTabSplitTests(SimpleTestCase):
    def test_headings_and_table_rows_stay_with_their_tab(self):
        markdown = """
Batangas State University
College of Informatics and Computing Sciences

# University Vision
A premier national university.

# Course Description
This course covers systems administration.

# Intended Learning Outcomes
1. Install an operating system.

# Table of Specifications
| Topic | Hours | Remember |
| Introduction | 2 | 40 |
| Storage | 4 | 10 |

# Course Policies
Attendance is required.
"""
        tabs = {tab['key']: tab['text'] for tab in split_markdown_into_tabs(markdown)}
        self.assertIn('premier national university', tabs['basic'].lower())
        self.assertIn('systems administration', tabs['details'].lower())
        self.assertIn('install an operating system', tabs['iloat'].lower())
        self.assertIn('Introduction | 2 | 40', tabs['tos'])
        self.assertIn('Attendance is required', tabs['policies'])
        self.assertNotIn('Introduction | 2 | 40', tabs.get('schedule', ''))

    def test_parent_section_does_not_keep_later_tabs(self):
        markdown = """
# Course Information
Batangas State University
College of Informatics and Computing Sciences

## Course Description
This course covers systems administration.

## Intended Learning Outcomes
Install an operating system.

## Course Learning Plan
Week 1 | Orientation

## Grading System
Excellent | 1.00 | 98 - 100

## Course Policies
Attendance is required.
"""
        tabs = {tab['key']: tab['text'] for tab in split_markdown_into_tabs(markdown)}
        self.assertIn('systems administration', tabs['details'].lower())
        self.assertIn('install an operating system', tabs['iloat'].lower())
        self.assertIn('Week 1 | Orientation', tabs['schedule'])
        self.assertIn('Excellent | 1.00', tabs['grading'])
        self.assertIn('Attendance is required', tabs['policies'])
        self.assertNotIn('systems administration', tabs['basic'].lower())

    def test_course_content_and_outcomes_are_not_left_in_basic(self):
        markdown = """
# Course Information
Batangas State University

# Course Learning Outcomes
Configure a network service.

# Course Content
Week 1 | Network models
"""
        tabs = {tab['key']: tab['text'] for tab in split_markdown_into_tabs(markdown)}
        self.assertIn('configure a network service', tabs['iloat'].lower())
        self.assertIn('Week 1 | Network models', tabs['schedule'])
        self.assertNotIn('network service', tabs['basic'].lower())

    def test_rewritten_sentences_are_dropped_and_missing_text_is_kept(self):
        source = (
            'This course covers systems administration.\n'
            'Install an operating system.'
        )
        parsed = {
            'description': 'This course covers systems administration. It also teaches cooking.',
            'ilos': [
                'Install an operating system.',
                'Invented outcome that is not written in the file.',
            ],
            'vision': 'A completely different vision statement for the university.',
            'passing_grade': 99,
        }
        cleaned = faithful_parsed_data(parsed, source)
        self.assertEqual(cleaned['description'], 'This course covers systems administration.')
        self.assertEqual(cleaned['ilos'], ['Install an operating system.'])
        self.assertEqual(cleaned['vision'], '')
        self.assertEqual(cleaned['passing_grade'], '')
        wrapped = faithful_parsed_data(
            {'description': 'This course covers systems administration.'},
            'This course covers\nsystems administration.',
        )
        self.assertEqual(wrapped['description'], 'This course covers systems administration.')

        merged = supplement_markdown(
            'Course Description\nShort.',
            'Course Description\nShort.\n\nStudents configure servers in the laboratory each week.',
        )
        self.assertIn('laboratory', merged)


class SyllabusJsonImportTests(SimpleTestCase):
    def test_file_text_becomes_fields_and_extra_sections(self):
        markdown = """
# Course Information
Course Code: IT 101
Course Title: Systems Administration

# Course Description
This course covers systems administration.

# Intended Learning Outcomes
1. Install an operating system.
2. Configure a network service.

# Course Learning Plan
| Week | Topic | Outcomes |
| 1 | Orientation | Install an OS |
| 2 | Storage | Partition a disk |

# Grading System
| Rating | Grade | Range |
| Excellent | 1.00 | 98-100 |

# Laboratory Safety
Wear closed shoes in the laboratory.
"""
        content = syllabus_content_from_markdown(markdown)
        self.assertEqual(content['course_code'], 'IT 101')
        self.assertEqual(content['course_title'], 'Systems Administration')
        self.assertIn('systems administration', content['description'].lower())
        self.assertEqual(content['ilos'], [
            'Install an operating system.',
            'Configure a network service.',
        ])
        self.assertEqual(len(content['tla_chapters']), 2)
        self.assertEqual(content['tla_chapters'][0]['topics'], 'Orientation')
        self.assertEqual(content['tla_chapters'][0]['ilo'], 'Install an OS')
        self.assertEqual(content['grading_rows'], [{'rating': 'Excellent', 'grade': '1.00', 'range': '98-100'}])
        self.assertNotIn('imported_sections', content)
        self.assertEqual(content['policies'][0]['title'], 'Laboratory Safety')
        self.assertIn('closed shoes', content['policies'][0]['content'])

    def test_table_cells_fill_the_existing_sections(self):
        markdown = """
| Course Code | Course Code | IT 101 |
| --- | --- | --- |
| Course Title | Systems Administration |
| Course Description | This course covers systems administration. |
| Intended Learning Outcomes |  |
| 1. Install an operating system. |  |
| 2. Configure a network service. |  |
| Week | Topic | Outcomes |
| --- | --- | --- |
| 1 | Orientation | Install an OS |
"""
        content = syllabus_content_from_markdown(markdown)
        self.assertEqual(content['course_code'], 'IT 101')
        self.assertEqual(content['course_title'], 'Systems Administration')
        self.assertIn('systems administration', content['description'].lower())
        self.assertEqual(content['ilos'], [
            'Install an operating system.',
            'Configure a network service.',
        ])
        self.assertEqual(content['tla_chapters'][0]['topics'], 'Orientation')
        self.assertEqual(content['tla_chapters'][0]['ilo'], 'Install an OS')


class TosExportDocumentTests(SimpleTestCase):
    """The standalone ToS download follows the institutional form."""

    def _text(self, document):
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return '\n'.join(parts)

    def test_form_names_the_examination_and_uses_course_fields(self):
        exams = [
            {'examLabel': 'Prelim', 'preparedBy': 'Prof. One'},
            {'examLabel': 'Midterm Examination', 'preparedBy': 'Prof. Two'},
            {'examLabel': 'Finals', 'preparedBy': 'Prof. Three'},
        ]
        data = {
            'course_code': 'IT 321',
            'course_title': 'Analytics Techniques and Tools',
            'semester_year': 'First Semester, AY 2026 - 2027',
            # Stale course values must not reach the form's letterhead.
            'college': 'SOME OTHER COLLEGE',
            'program': 'BatStateU-Lipa Campus',
            'document_title': 'JUNK DOCUMENT TITLE',
            'tos': [
                {
                    **exam,
                    'topics': [
                        {
                            'topic': 'Predictive Analytics',
                            'ilos': ['ILO1'],
                            'hours_class': 8,
                            'hours_min': 72,
                            'weight': 50,
                            'bloom_items': {'remember': 5, 'understand': 10},
                            'bloom_pct': {'remember': 33.3, 'understand': 66.7},
                            'total_points': 50,
                        }
                    ],
                }
                for exam in exams
            ],
        }

        text = self._text(build_tos_docx(data))

        for title in ('Preliminary Examination', 'Midterm Examination', 'Final Examination'):
            self.assertIn(title, text)
        self.assertIn('TABLE OF SPECIFICATIONS', text)
        self.assertIn('COURSE CODE : IT 321', text)
        self.assertIn('COURSE TITLE: Analytics Techniques and Tools', text)
        self.assertIn('Prof. One', text)
        self.assertIn('Predictive Analytics', text)
        self.assertIn('Prepare', text)
        self.assertIn('Checked and Verified by:', text)
        self.assertIn('Approved by:', text)
        self.assertIn('*ILO - Intended Learning Outcomes', text)
        self.assertNotIn('SOME OTHER COLLEGE', text)
        self.assertNotIn('BatStateU-Lipa Campus', text)
        self.assertNotIn('JUNK DOCUMENT TITLE', text)


class AssessmentDocxFormatTests(SimpleTestCase):
    """A downloaded assessment is the exam, then answer letters on the last page."""

    def _paragraphs(self, document):
        return [paragraph.text for paragraph in document.paragraphs]

    def _has_page_break(self, document):
        from docx.oxml.ns import qn

        for paragraph in document.paragraphs:
            for br in paragraph._element.findall('.//' + qn('w:br')):
                if br.get(qn('w:type')) == 'page':
                    return True
        return False

    def test_questions_and_choices_then_answer_letters_on_the_last_page(self):
        document = _build_plain_assessment_docx(
            {
                'title': 'IT 321 Prelim Examination',
                'assessment_type': 'prelim',
                'total_items': 2,
                'items': [
                    {
                        'item_number': 1,
                        'section_title': 'Foundations of HCI',
                        'bloom_level': 'Remember',
                        'bloom_verb': 'Identify',
                        'ilo_refs': ['ILO1'],
                        'question_text': 'Which field studies how people use computers?',
                        'choices': ['A. Human-computer interaction', 'Graphic design', 'Accounting', 'Botany'],
                        'correct_answer': 'A',
                        'answer_key': 'HCI is the study of people and computers.',
                    },
                    {
                        'item_number': 2,
                        'question_text': 'A prototype is used to do what?',
                        'choices': ['Hide the design', 'B. Try an idea with users', 'Skip testing', 'Replace the users'],
                        'correct_answer': 'B',
                        'answer_key': 'Prototypes are tried with users.',
                    },
                ],
            },
            include_choices=True,
            include_answer_key=True,
        )

        paragraphs = self._paragraphs(document)
        text = '\n'.join(paragraphs)
        answer_at = paragraphs.index('1. A')

        self.assertIn('IT 321 Prelim Examination', text)
        self.assertIn('1. Which field studies how people use computers?', text)
        self.assertIn('A. Human-computer interaction', text)
        self.assertIn('B. Graphic design', text)
        self.assertIn('2. A prototype is used to do what?', text)
        self.assertIn('B. Try an idea with users', text)
        self.assertEqual(paragraphs[answer_at:answer_at + 2], ['1. A', '2. B'])
        self.assertTrue(self._has_page_break(document))
        self.assertLess(paragraphs.index('1. Which field studies how people use computers?'), answer_at)
        for label in (
            'Bloom',
            'Remember',
            'Identify',
            'ILO',
            'Foundations of HCI',
            'Correct Answer',
            'Answer Key',
            'Explanation',
            'Type:',
            'Total Items',
            'HCI is the study of people and computers.',
        ):
            self.assertNotIn(label, text)

    def test_exam_copy_omits_the_answer_page(self):
        document = _build_plain_assessment_docx(
            {
                'title': 'Quiz',
                'items': [
                    {
                        'item_number': 1,
                        'question_text': 'Pick one.',
                        'choices': ['First', 'Second'],
                        'correct_answer': 'A',
                        'answer_key': 'Because the first choice is right.',
                    },
                ],
            },
            include_choices=True,
            include_answer_key=False,
        )
        text = '\n'.join(self._paragraphs(document))
        self.assertIn('1. Pick one.', text)
        self.assertIn('A. First', text)
        self.assertNotIn('1. A', text)
        self.assertNotIn('Because the first choice is right.', text)
        self.assertFalse(self._has_page_break(document))
